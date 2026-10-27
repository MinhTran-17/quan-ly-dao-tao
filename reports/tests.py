from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase, tag
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from academics.models import HocPhan, LopHanhChinh, NganhHoc
from accounts.models import VaiTro
from classes.models import LopHocPhan
from enrollment.models import DangKyHocPhan
from grades.models import Diem
from people.models import CoVanHocTap, GiangVien, SinhVien

TaiKhoan = get_user_model()
MAT_KHAU = 'MatKhauManh123'


def tao_sinh_vien(ma, ma_lop):
    tk = TaiKhoan.objects.create_user(ten_dang_nhap=f'tk_{ma}', vai_tro=VaiTro.SINH_VIEN, password=MAT_KHAU)
    return SinhVien.objects.create(
        ma_sinh_vien=ma, ho_ten=f'SV {ma}', ngay_sinh='2004-01-01', email=f'{ma}@test.local',
        so_dien_thoai='0900000000', ma_lop=ma_lop, ma_tai_khoan=tk,
    )


def tao_giang_vien(ma):
    tk = TaiKhoan.objects.create_user(ten_dang_nhap=f'tk_{ma}', vai_tro=VaiTro.GIANG_VIEN, password=MAT_KHAU)
    return GiangVien.objects.create(
        ma_giang_vien=ma, ho_ten=f'GV {ma}', email=f'{ma}@test.local', so_dien_thoai='0900000000', ma_tai_khoan=tk,
    )


def tao_lop_hanh_chinh(ma='CNTT1'):
    nganh, _ = NganhHoc.objects.get_or_create(ma_nganh='CNTT', defaults={'ten_nganh': 'Công nghệ thông tin'})
    tk = TaiKhoan.objects.create_user(ten_dang_nhap=f'tk_cv_{ma}', vai_tro=VaiTro.CO_VAN_HOC_TAP, password=MAT_KHAU)
    co_van = CoVanHocTap.objects.create(
        ma_co_van=f'CV_{ma}', ho_ten='Cố vấn', email='cv@test.local', so_dien_thoai='0900000000', ma_tai_khoan=tk,
    )
    return LopHanhChinh.objects.create(ma_lop=ma, ten_lop=f'Lớp {ma}', ma_nganh=nganh, ma_co_van=co_van)


@tag('tich_hop')
class ThongKeReportTest(TestCase):
    def setUp(self):
        self.phong_dao_tao = TaiKhoan.objects.create_user(
            ten_dang_nhap='pdt', vai_tro=VaiTro.PHONG_DAO_TAO, password=MAT_KHAU
        )
        self.client.force_login(self.phong_dao_tao)

        lop_hanh_chinh = tao_lop_hanh_chinh('CNTT1')
        giang_vien = tao_giang_vien('GV01')

        self.csdl = HocPhan.objects.create(ma_hoc_phan='CSDL', ten_hoc_phan='Cơ sở dữ liệu', so_tin_chi=3)
        self.ctdl = HocPhan.objects.create(ma_hoc_phan='CTDL', ten_hoc_phan='Cấu trúc dữ liệu', so_tin_chi=3)

        self.lop_csdl = LopHocPhan.objects.create(
            ma_lop_hoc_phan='CSDL_01', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=2,
            ma_hoc_phan=self.csdl, ma_giang_vien=giang_vien,
        )
        self.lop_ctdl = LopHocPhan.objects.create(
            ma_lop_hoc_phan='CTDL_01', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=5,
            ma_hoc_phan=self.ctdl, ma_giang_vien=giang_vien,
        )
        # Lớp học kỳ khác để kiểm tra bộ lọc không lẫn dữ liệu
        self.lop_csdl_hk2 = LopHocPhan.objects.create(
            ma_lop_hoc_phan='CSDL_02', hoc_ky='2', nam_hoc='2025-2026', si_so_toi_da=10,
            ma_hoc_phan=self.csdl, ma_giang_vien=giang_vien,
        )

        sv1 = tao_sinh_vien('SV001', lop_hanh_chinh)
        sv2 = tao_sinh_vien('SV002', lop_hanh_chinh)
        sv3 = tao_sinh_vien('SV003', lop_hanh_chinh)

        dk1 = DangKyHocPhan.objects.create(
            ma_sinh_vien=sv1, ma_lop_hoc_phan=self.lop_csdl, thoi_gian_dang_ky=timezone.now()
        )
        dk2 = DangKyHocPhan.objects.create(
            ma_sinh_vien=sv2, ma_lop_hoc_phan=self.lop_csdl, thoi_gian_dang_ky=timezone.now()
        )
        dk3 = DangKyHocPhan.objects.create(
            ma_sinh_vien=sv3, ma_lop_hoc_phan=self.lop_ctdl, thoi_gian_dang_ky=timezone.now()
        )

        Diem.objects.create(ma_dang_ky=dk1, diem_qua_trinh=Decimal('8'), diem_thi=Decimal('8'))  # 8.0 - Giỏi, đạt
        Diem.objects.create(ma_dang_ky=dk2, diem_qua_trinh=Decimal('2'), diem_thi=Decimal('2'))  # 2.0 - Không đạt
        Diem.objects.create(ma_dang_ky=dk3, diem_qua_trinh=Decimal('9'), diem_thi=Decimal('9'))  # 9.0 - Xuất sắc, đạt

    def test_si_so_theo_lop_dung_so_lieu(self):
        response = self.client.get(reverse('reports:si_so_theo_lop'))
        self.assertEqual(response.status_code, 200)
        theo_ma = {lop.ma_lop_hoc_phan: lop.so_dang_ky for lop in response.context['danh_sach']}
        self.assertEqual(theo_ma['CSDL_01'], 2)
        self.assertEqual(theo_ma['CTDL_01'], 1)
        self.assertEqual(theo_ma['CSDL_02'], 0)

    def test_si_so_theo_lop_loc_theo_hoc_ky(self):
        response = self.client.get(reverse('reports:si_so_theo_lop'), {'hoc_ky': '2'})
        ma_cac_lop = [lop.ma_lop_hoc_phan for lop in response.context['danh_sach']]
        self.assertEqual(ma_cac_lop, ['CSDL_02'])

    def test_hoc_phan_nhieu_dang_ky_xep_hang_dung(self):
        response = self.client.get(reverse('reports:hoc_phan_nhieu_dang_ky'))
        danh_sach = list(response.context['danh_sach'])
        self.assertEqual(danh_sach[0].pk, 'CSDL')
        self.assertEqual(danh_sach[0].so_dang_ky, 2)
        self.assertEqual(danh_sach[1].pk, 'CTDL')
        self.assertEqual(danh_sach[1].so_dang_ky, 1)

    def test_hoc_phan_nhieu_dang_ky_loc_theo_hoc_ky(self):
        response = self.client.get(reverse('reports:hoc_phan_nhieu_dang_ky'), {'hoc_ky': '2'})
        theo_ma = {hp.pk: hp.so_dang_ky for hp in response.context['danh_sach']}
        self.assertEqual(theo_ma['CSDL'], 0)
        self.assertEqual(theo_ma['CTDL'], 0)

    def test_ket_qua_hoc_tap_phan_bo_va_ty_le_dat(self):
        response = self.client.get(reverse('reports:ket_qua_hoc_tap'))
        self.assertEqual(response.context['tong_so_diem'], 3)
        phan_bo = response.context['phan_bo']
        self.assertEqual(phan_bo['duoi_4'], 1)
        self.assertEqual(phan_bo['tu_7_den_duoi_8_5'], 1)
        self.assertEqual(phan_bo['tu_8_5_tro_len'], 1)
        self.assertEqual(phan_bo['tu_4_den_duoi_5_5'], 0)
        self.assertEqual(phan_bo['tu_5_5_den_duoi_7'], 0)

        theo_hoc_phan = {muc['ten']: muc for muc in response.context['theo_hoc_phan']}
        self.assertEqual(theo_hoc_phan['Cơ sở dữ liệu']['tong'], 2)
        self.assertEqual(theo_hoc_phan['Cơ sở dữ liệu']['dat'], 1)
        self.assertEqual(theo_hoc_phan['Cơ sở dữ liệu']['ty_le_dat'], 50.0)
        self.assertEqual(theo_hoc_phan['Cấu trúc dữ liệu']['ty_le_dat'], 100.0)

    def test_khong_phai_phong_dao_tao_bi_chan(self):
        giang_vien_khac = TaiKhoan.objects.create_user(
            ten_dang_nhap='gv_khac', vai_tro=VaiTro.GIANG_VIEN, password=MAT_KHAU
        )
        self.client.force_login(giang_vien_khac)
        for ten_url in ('reports:si_so_theo_lop', 'reports:hoc_phan_nhieu_dang_ky', 'reports:ket_qua_hoc_tap'):
            response = self.client.get(reverse(ten_url))
            self.assertEqual(response.status_code, 403)

    def test_chua_dang_nhap_bi_chuyen_huong(self):
        self.client.logout()
        response = self.client.get(reverse('reports:si_so_theo_lop'))
        self.assertEqual(response.status_code, 302)


@tag('tich_hop')
class XuatExcelTest(ThongKeReportTest):
    def test_xuat_excel_si_so_theo_lop_giu_bo_loc(self):
        response = self.client.get(reverse('reports:si_so_theo_lop_xuat_excel'), {'hoc_ky': '2'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        sheet = load_workbook(BytesIO(response.content)).active
        ma_cac_lop = [row[0].value for row in sheet.iter_rows(min_row=2)]
        self.assertEqual(ma_cac_lop, ['CSDL_02'])

    def test_xuat_excel_hoc_phan_nhieu_dang_ky_dung_thu_tu(self):
        response = self.client.get(reverse('reports:hoc_phan_nhieu_dang_ky_xuat_excel'))
        sheet = load_workbook(BytesIO(response.content)).active
        dong_dau = [cell.value for cell in sheet[2]]
        self.assertEqual(dong_dau[1], 'CSDL')
        self.assertEqual(dong_dau[3], 2)

    def test_xuat_excel_ket_qua_hoc_tap_co_hai_sheet(self):
        response = self.client.get(reverse('reports:ket_qua_hoc_tap_xuat_excel'))
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ['Phan bo diem', 'Ty le dat theo hoc phan'])
        sheet2 = workbook['Ty le dat theo hoc phan']
        cac_dong = {row[0].value: row for row in sheet2.iter_rows(min_row=2)}
        self.assertEqual(cac_dong['Cơ sở dữ liệu'][3].value, 50.0)

    def test_khong_phai_phong_dao_tao_bi_chan_xuat_excel(self):
        giang_vien_khac = TaiKhoan.objects.create_user(
            ten_dang_nhap='gv_khac_excel', vai_tro=VaiTro.GIANG_VIEN, password=MAT_KHAU
        )
        self.client.force_login(giang_vien_khac)
        response = self.client.get(reverse('reports:si_so_theo_lop_xuat_excel'))
        self.assertEqual(response.status_code, 403)


def kiem_tra_dinh_dang_excel(test, noi_dung_tep):
    """Mọi sheet của file xuất: tiêu đề in đậm, cố định dòng 1, cột giãn theo nội dung dài nhất (tối đa 60)."""
    from io import BytesIO as _BytesIO

    from openpyxl import load_workbook as _load_workbook
    from openpyxl.utils import get_column_letter as _cot
    workbook = _load_workbook(_BytesIO(noi_dung_tep))
    for sheet in workbook.worksheets:
        test.assertEqual(sheet.freeze_panes, 'A2', sheet.title)
        test.assertTrue(all(cell.font.bold for cell in sheet[1] if cell.value is not None), sheet.title)
        for chi_so, cot in enumerate(sheet.iter_cols(values_only=True), start=1):
            dai_nhat = max(len(str(v)) for v in cot if v is not None)
            test.assertEqual(sheet.column_dimensions[_cot(chi_so)].width, min(max(dai_nhat + 2, 8), 60), (sheet.title, chi_so))
    return workbook


@tag('tich_hop')
class DinhDangFileExcelBaoCaoTest(ThongKeReportTest):
    def test_ca_ba_bao_cao_xuat_excel_gian_cot_in_dam_co_dinh_tieu_de(self):
        for ten_url in ('reports:si_so_theo_lop_xuat_excel', 'reports:hoc_phan_nhieu_dang_ky_xuat_excel',
                        'reports:ket_qua_hoc_tap_xuat_excel'):
            response = self.client.get(reverse(ten_url))
            self.assertEqual(response.status_code, 200, ten_url)
            workbook = kiem_tra_dinh_dang_excel(self, response.content)
            self.assertGreaterEqual(len(workbook.worksheets), 1)
