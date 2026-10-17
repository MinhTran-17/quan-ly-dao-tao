import re
from datetime import date, time, timedelta
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from academics.models import DieuKienTienQuyet, DotDangKy, HocKy, HocPhan, LopHanhChinh, NganhHoc
from accounts.models import VaiTro
from accounts.tests import kiem_tra_can_bang_the
from classes.models import LichHoc, LichThi, LopHocPhan
from grades.models import Diem
from people.models import CoVanHocTap, GiangVien, SinhVien

from .models import DangKyHocPhan, TrangThaiDangKy
from .services import LoiDangKyHocPhan, dang_ky_hoc_phan, huy_dang_ky_hoc_phan

TaiKhoan = get_user_model()
MAT_KHAU = 'MatKhauManh123'


def tao_dot_dang_ky_mo(hoc_ky_so, nam_hoc):
    """Học kỳ trong danh mục + một đợt đăng ký đang mở ngay bây giờ — fixture dùng chung
    cho các test gọi dang_ky_hoc_phan()/xac_dinh_trang_thai() (nay phải qua kiểm tra đợt
    đăng ký trước, xem enrollment/services.py)."""
    hoc_ky_obj = HocKy.objects.create(
        hoc_ky=int(hoc_ky_so), nam_hoc=nam_hoc,
        ngay_bat_dau=timezone.localdate(), ngay_ket_thuc=timezone.localdate() + timedelta(days=120),
    )
    return DotDangKy.objects.create(
        hoc_ky=hoc_ky_obj, ngay_mo=timezone.now() - timedelta(days=1), ngay_dong=timezone.now() + timedelta(days=30),
    )


def tao_sinh_vien(ma_sinh_vien, ma_lop):
    tk = TaiKhoan.objects.create_user(
        ten_dang_nhap=f'tk_{ma_sinh_vien}', vai_tro=VaiTro.SINH_VIEN, password=MAT_KHAU
    )
    return SinhVien.objects.create(
        ma_sinh_vien=ma_sinh_vien, ho_ten=f'Sinh viên {ma_sinh_vien}', ngay_sinh='2004-01-01',
        email=f'{ma_sinh_vien}@test.local', so_dien_thoai='0900000000', ma_lop=ma_lop, ma_tai_khoan=tk,
    )


def tao_giang_vien(ma_giang_vien):
    tk = TaiKhoan.objects.create_user(
        ten_dang_nhap=f'tk_{ma_giang_vien}', vai_tro=VaiTro.GIANG_VIEN, password=MAT_KHAU
    )
    return GiangVien.objects.create(
        ma_giang_vien=ma_giang_vien, ho_ten=f'Giảng viên {ma_giang_vien}', email=f'{ma_giang_vien}@test.local',
        so_dien_thoai='0900000000', ma_tai_khoan=tk,
    )


def tao_lop_hanh_chinh(ma_lop='CNTT1'):
    nganh, _ = NganhHoc.objects.get_or_create(ma_nganh='CNTT', defaults={'ten_nganh': 'Công nghệ thông tin'})
    tk_cv = TaiKhoan.objects.create_user(
        ten_dang_nhap=f'tk_cv_{ma_lop}', vai_tro=VaiTro.CO_VAN_HOC_TAP, password=MAT_KHAU
    )
    co_van = CoVanHocTap.objects.create(
        ma_co_van=f'CV_{ma_lop}', ho_ten='Cố vấn Test', email='cv@test.local',
        so_dien_thoai='0900000000', ma_tai_khoan=tk_cv,
    )
    return LopHanhChinh.objects.create(ma_lop=ma_lop, ten_lop=f'Lớp {ma_lop}', ma_nganh=nganh, ma_co_van=co_van)


class DangKyHocPhanServiceTest(TestCase):
    def setUp(self):
        tao_dot_dang_ky_mo('1', '2025-2026')
        self.lop_hanh_chinh = tao_lop_hanh_chinh('CNTT1')
        self.sinh_vien = tao_sinh_vien('SV001', self.lop_hanh_chinh)
        self.giang_vien = tao_giang_vien('GV001')
        self.ctdl = HocPhan.objects.create(ma_hoc_phan='CTDL', ten_hoc_phan='Cấu trúc dữ liệu', so_tin_chi=3)
        self.csdl = HocPhan.objects.create(ma_hoc_phan='CSDL', ten_hoc_phan='Cơ sở dữ liệu', so_tin_chi=3)
        DieuKienTienQuyet.objects.create(ma_hoc_phan=self.csdl, ma_hoc_phan_tien_quyet=self.ctdl)
        self.lop_csdl = LopHocPhan.objects.create(
            ma_lop_hoc_phan='CSDL_01', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=2,
            ma_hoc_phan=self.csdl, ma_giang_vien=self.giang_vien,
        )
        LichHoc.objects.create(ma_lop_hoc_phan=self.lop_csdl, thu=2, tiet_bat_dau=1, so_tiet=3, phong_hoc='A101')

    def _cho_dat_ctdl(self, sinh_vien, ma_lop_hoc_phan_cu):
        """Tạo một lần đăng ký CTDL đã đạt cho sinh_vien, để test trùng lịch/sĩ số
        không bị chặn nhầm bởi điều kiện tiên quyết của CSDL."""
        lop_ctdl_cu = LopHocPhan.objects.create(
            ma_lop_hoc_phan=ma_lop_hoc_phan_cu, hoc_ky='1', nam_hoc='2024-2025', si_so_toi_da=40,
            ma_hoc_phan=self.ctdl, ma_giang_vien=self.giang_vien,
        )
        dang_ky_cu = DangKyHocPhan.objects.create(
            ma_sinh_vien=sinh_vien, ma_lop_hoc_phan=lop_ctdl_cu, thoi_gian_dang_ky=timezone.now(),
        )
        Diem.objects.create(ma_dang_ky=dang_ky_cu, diem_qua_trinh=8, diem_thi=8)

    def test_dang_ky_hop_le_khi_khong_co_tien_quyet(self):
        lop_khong_tien_quyet = LopHocPhan.objects.create(
            ma_lop_hoc_phan='CTDL_01', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=40,
            ma_hoc_phan=self.ctdl, ma_giang_vien=self.giang_vien,
        )
        dang_ky = dang_ky_hoc_phan(self.sinh_vien, lop_khong_tien_quyet)
        self.assertTrue(DangKyHocPhan.objects.filter(pk=dang_ky.pk).exists())

    def test_thieu_hoc_phan_tien_quyet(self):
        with self.assertRaises(LoiDangKyHocPhan) as ngu_canh:
            dang_ky_hoc_phan(self.sinh_vien, self.lop_csdl)
        self.assertIn('tiên quyết', str(ngu_canh.exception))
        self.assertIn('Cấu trúc dữ liệu', str(ngu_canh.exception))
        self.assertFalse(DangKyHocPhan.objects.filter(ma_lop_hoc_phan=self.lop_csdl).exists())

    def test_dang_ky_duoc_khi_da_dat_tien_quyet(self):
        lop_ctdl_cu = LopHocPhan.objects.create(
            ma_lop_hoc_phan='CTDL_CU', hoc_ky='1', nam_hoc='2024-2025', si_so_toi_da=40,
            ma_hoc_phan=self.ctdl, ma_giang_vien=self.giang_vien,
        )
        dang_ky_cu = DangKyHocPhan.objects.create(
            ma_sinh_vien=self.sinh_vien, ma_lop_hoc_phan=lop_ctdl_cu, thoi_gian_dang_ky=timezone.now(),
        )
        Diem.objects.create(ma_dang_ky=dang_ky_cu, diem_qua_trinh=5, diem_thi=5)

        dang_ky_moi = dang_ky_hoc_phan(self.sinh_vien, self.lop_csdl)
        self.assertTrue(DangKyHocPhan.objects.filter(pk=dang_ky_moi.pk).exists())

    def test_chua_du_diem_thi_thi_chua_tinh_la_dat(self):
        lop_ctdl_cu = LopHocPhan.objects.create(
            ma_lop_hoc_phan='CTDL_CU2', hoc_ky='1', nam_hoc='2024-2025', si_so_toi_da=40,
            ma_hoc_phan=self.ctdl, ma_giang_vien=self.giang_vien,
        )
        dang_ky_cu = DangKyHocPhan.objects.create(
            ma_sinh_vien=self.sinh_vien, ma_lop_hoc_phan=lop_ctdl_cu, thoi_gian_dang_ky=timezone.now(),
        )
        Diem.objects.create(ma_dang_ky=dang_ky_cu, diem_qua_trinh=9, diem_thi=None)

        with self.assertRaises(LoiDangKyHocPhan):
            dang_ky_hoc_phan(self.sinh_vien, self.lop_csdl)

    def test_trung_lich_cung_hoc_ky_nam_hoc(self):
        lop_khac = LopHocPhan.objects.create(
            ma_lop_hoc_phan='KHAC_01', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=40,
            ma_hoc_phan=self.ctdl, ma_giang_vien=self.giang_vien,
        )
        LichHoc.objects.create(ma_lop_hoc_phan=lop_khac, thu=2, tiet_bat_dau=2, so_tiet=3, phong_hoc='B101')
        DangKyHocPhan.objects.create(
            ma_sinh_vien=self.sinh_vien, ma_lop_hoc_phan=lop_khac, thoi_gian_dang_ky=timezone.now(),
        )
        Diem.objects.create(
            ma_dang_ky=DangKyHocPhan.objects.get(ma_sinh_vien=self.sinh_vien, ma_lop_hoc_phan=lop_khac),
            diem_qua_trinh=8, diem_thi=8,
        )

        with self.assertRaises(LoiDangKyHocPhan) as ngu_canh:
            dang_ky_hoc_phan(self.sinh_vien, self.lop_csdl)
        self.assertIn('trùng', str(ngu_canh.exception))

    def test_khong_trung_lich_neu_khac_hoc_ky(self):
        self._cho_dat_ctdl(self.sinh_vien, 'CTDL_CU_1')
        lop_khac = LopHocPhan.objects.create(
            ma_lop_hoc_phan='KHAC_02', hoc_ky='2', nam_hoc='2025-2026', si_so_toi_da=40,
            ma_hoc_phan=self.ctdl, ma_giang_vien=self.giang_vien,
        )
        LichHoc.objects.create(ma_lop_hoc_phan=lop_khac, thu=2, tiet_bat_dau=1, so_tiet=3, phong_hoc='B101')
        DangKyHocPhan.objects.create(
            ma_sinh_vien=self.sinh_vien, ma_lop_hoc_phan=lop_khac, thoi_gian_dang_ky=timezone.now(),
        )
        dang_ky_hoc_phan(self.sinh_vien, self.lop_csdl)
        self.assertTrue(DangKyHocPhan.objects.filter(ma_lop_hoc_phan=self.lop_csdl).exists())

    def test_lop_da_du_si_so(self):
        sv2 = tao_sinh_vien('SV002', self.lop_hanh_chinh)
        DangKyHocPhan.objects.create(
            ma_sinh_vien=self.sinh_vien, ma_lop_hoc_phan=self.lop_csdl, thoi_gian_dang_ky=timezone.now(),
        )
        DangKyHocPhan.objects.create(
            ma_sinh_vien=sv2, ma_lop_hoc_phan=self.lop_csdl, thoi_gian_dang_ky=timezone.now(),
        )
        sv3 = tao_sinh_vien('SV003', self.lop_hanh_chinh)
        self._cho_dat_ctdl(sv3, 'CTDL_CU_SV3')
        with self.assertRaises(LoiDangKyHocPhan) as ngu_canh:
            dang_ky_hoc_phan(sv3, self.lop_csdl)
        self.assertIn('sĩ số', str(ngu_canh.exception))

    def test_dang_ky_trung(self):
        DangKyHocPhan.objects.create(
            ma_sinh_vien=self.sinh_vien, ma_lop_hoc_phan=self.lop_csdl, thoi_gian_dang_ky=timezone.now(),
        )
        with self.assertRaises(LoiDangKyHocPhan) as ngu_canh:
            dang_ky_hoc_phan(self.sinh_vien, self.lop_csdl)
        self.assertIn('đã đăng ký', str(ngu_canh.exception))

    def test_uu_tien_kiem_tra_tien_quyet_truoc_trung_lich(self):
        # sinh viên chưa đạt CTDL (tiên quyết của CSDL) và cũng đăng ký một lớp trùng lịch.
        lop_khac = LopHocPhan.objects.create(
            ma_lop_hoc_phan='KHAC_03', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=40,
            ma_hoc_phan=self.ctdl, ma_giang_vien=self.giang_vien,
        )
        LichHoc.objects.create(ma_lop_hoc_phan=lop_khac, thu=2, tiet_bat_dau=1, so_tiet=3, phong_hoc='B101')
        DangKyHocPhan.objects.create(
            ma_sinh_vien=self.sinh_vien, ma_lop_hoc_phan=lop_khac, thoi_gian_dang_ky=timezone.now(),
        )
        with self.assertRaises(LoiDangKyHocPhan) as ngu_canh:
            dang_ky_hoc_phan(self.sinh_vien, self.lop_csdl)
        self.assertIn('tiên quyết', str(ngu_canh.exception))


class QuyenDangKyViewTest(TestCase):
    def setUp(self):
        self.lop_hanh_chinh = tao_lop_hanh_chinh('CNTT1')
        self.sinh_vien = tao_sinh_vien('SV010', self.lop_hanh_chinh)
        self.giang_vien = tao_giang_vien('GV010')
        self.giang_vien_khac = tao_giang_vien('GV011')
        self.phong_dao_tao = TaiKhoan.objects.create_user(
            ten_dang_nhap='pdt', vai_tro=VaiTro.PHONG_DAO_TAO, password=MAT_KHAU
        )
        self.hoc_phan = HocPhan.objects.create(ma_hoc_phan='CTDL', ten_hoc_phan='Cấu trúc dữ liệu', so_tin_chi=3)
        self.lop_hoc_phan = LopHocPhan.objects.create(
            ma_lop_hoc_phan='CTDL_01', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=40,
            ma_hoc_phan=self.hoc_phan, ma_giang_vien=self.giang_vien,
        )

    def test_sinh_vien_xem_duoc_lop_dang_mo(self):
        self.client.force_login(self.sinh_vien.ma_tai_khoan)
        response = self.client.get(reverse('enrollment:lop_dang_mo'))
        self.assertEqual(response.status_code, 200)

    def test_giang_vien_khong_xem_duoc_lop_dang_mo(self):
        self.client.force_login(self.giang_vien.ma_tai_khoan)
        response = self.client.get(reverse('enrollment:lop_dang_mo'))
        self.assertEqual(response.status_code, 403)

    def test_giang_vien_xem_duoc_sinh_vien_lop_minh(self):
        self.client.force_login(self.giang_vien.ma_tai_khoan)
        response = self.client.get(
            reverse('enrollment:danh_sach_sinh_vien_lop', args=[self.lop_hoc_phan.pk])
        )
        self.assertEqual(response.status_code, 200)

    def test_giang_vien_khong_xem_duoc_sinh_vien_lop_khac(self):
        self.client.force_login(self.giang_vien_khac.ma_tai_khoan)
        response = self.client.get(
            reverse('enrollment:danh_sach_sinh_vien_lop', args=[self.lop_hoc_phan.pk])
        )
        self.assertEqual(response.status_code, 403)

    def test_phong_dao_tao_xem_duoc_bat_ky_lop_nao(self):
        self.client.force_login(self.phong_dao_tao)
        response = self.client.get(
            reverse('enrollment:danh_sach_sinh_vien_lop', args=[self.lop_hoc_phan.pk])
        )
        self.assertEqual(response.status_code, 200)

    def test_chua_dang_nhap_bi_chuyen_huong(self):
        response = self.client.get(reverse('enrollment:lop_dang_mo'))
        self.assertEqual(response.status_code, 302)


class DangKyQuaViewTest(TestCase):
    def setUp(self):
        tao_dot_dang_ky_mo('1', '2025-2026')
        self.lop_hanh_chinh = tao_lop_hanh_chinh('CNTT1')
        self.sinh_vien = tao_sinh_vien('SV020', self.lop_hanh_chinh)
        self.giang_vien = tao_giang_vien('GV020')
        self.hoc_phan = HocPhan.objects.create(ma_hoc_phan='CTDL', ten_hoc_phan='Cấu trúc dữ liệu', so_tin_chi=3)
        self.lop_hoc_phan = LopHocPhan.objects.create(
            ma_lop_hoc_phan='CTDL_01', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=40,
            ma_hoc_phan=self.hoc_phan, ma_giang_vien=self.giang_vien,
        )
        self.client.force_login(self.sinh_vien.ma_tai_khoan)

    def test_dang_ky_qua_view_thanh_cong(self):
        response = self.client.post(reverse('enrollment:dang_ky', args=[self.lop_hoc_phan.pk]))
        self.assertRedirects(response, reverse('enrollment:lop_dang_mo'))
        self.assertTrue(
            DangKyHocPhan.objects.filter(ma_sinh_vien=self.sinh_vien, ma_lop_hoc_phan=self.lop_hoc_phan).exists()
        )


class DotDangKyGioiHanTest(TestCase):
    """Phần A: chỉ cho đăng ký khi thời điểm hiện tại nằm trong một đợt đăng ký còn hiệu
    lực của đúng học kỳ/năm học của lớp — kiểm tra trước tiên quyết/trùng lịch/sĩ số."""

    def setUp(self):
        self.lop_hanh_chinh = tao_lop_hanh_chinh('CNTT1')
        self.sinh_vien = tao_sinh_vien('SV030', self.lop_hanh_chinh)
        self.giang_vien = tao_giang_vien('GV030')
        self.hoc_phan = HocPhan.objects.create(ma_hoc_phan='CTDL', ten_hoc_phan='Cấu trúc dữ liệu', so_tin_chi=3)
        self.lop_hoc_phan = LopHocPhan.objects.create(
            ma_lop_hoc_phan='CTDL_01', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=40,
            ma_hoc_phan=self.hoc_phan, ma_giang_vien=self.giang_vien,
        )

    def test_chan_khi_chua_co_dot_dang_ky_nao(self):
        with self.assertRaises(LoiDangKyHocPhan) as ngu_canh:
            dang_ky_hoc_phan(self.sinh_vien, self.lop_hoc_phan)
        self.assertIn('Chưa có đợt đăng ký', str(ngu_canh.exception))
        self.assertFalse(DangKyHocPhan.objects.filter(ma_lop_hoc_phan=self.lop_hoc_phan).exists())

    def test_chan_khi_dot_chua_mo(self):
        hoc_ky_obj = HocKy.objects.create(
            hoc_ky=1, nam_hoc='2025-2026', ngay_bat_dau=timezone.localdate(),
            ngay_ket_thuc=timezone.localdate() + timedelta(days=120),
        )
        DotDangKy.objects.create(
            hoc_ky=hoc_ky_obj, ngay_mo=timezone.now() + timedelta(days=5), ngay_dong=timezone.now() + timedelta(days=10),
        )
        with self.assertRaises(LoiDangKyHocPhan) as ngu_canh:
            dang_ky_hoc_phan(self.sinh_vien, self.lop_hoc_phan)
        self.assertIn('chưa mở', str(ngu_canh.exception))

    def test_chan_khi_dot_da_dong(self):
        hoc_ky_obj = HocKy.objects.create(
            hoc_ky=1, nam_hoc='2025-2026', ngay_bat_dau=timezone.localdate(),
            ngay_ket_thuc=timezone.localdate() + timedelta(days=120),
        )
        DotDangKy.objects.create(
            hoc_ky=hoc_ky_obj, ngay_mo=timezone.now() - timedelta(days=10), ngay_dong=timezone.now() - timedelta(days=5),
        )
        with self.assertRaises(LoiDangKyHocPhan) as ngu_canh:
            dang_ky_hoc_phan(self.sinh_vien, self.lop_hoc_phan)
        self.assertIn('đã đóng', str(ngu_canh.exception))

    def test_dang_ky_duoc_khi_dot_dang_mo(self):
        tao_dot_dang_ky_mo('1', '2025-2026')
        dang_ky = dang_ky_hoc_phan(self.sinh_vien, self.lop_hoc_phan)
        self.assertTrue(DangKyHocPhan.objects.filter(pk=dang_ky.pk).exists())

    def test_trang_web_bao_ro_ly_do_ngoai_dot(self):
        self.client.force_login(self.sinh_vien.ma_tai_khoan)
        response = self.client.post(reverse('enrollment:dang_ky', args=[self.lop_hoc_phan.pk]))
        self.assertRedirects(response, reverse('enrollment:lop_dang_mo'))
        self.assertFalse(DangKyHocPhan.objects.filter(ma_lop_hoc_phan=self.lop_hoc_phan).exists())
        response = self.client.get(reverse('enrollment:lop_dang_mo'))
        self.assertContains(response, 'Ngoài đợt đăng ký')


class HuyDangKyTest(TestCase):
    def setUp(self):
        tao_dot_dang_ky_mo('1', '2025-2026')
        self.lop_hanh_chinh = tao_lop_hanh_chinh('CNTT1')
        self.sinh_vien = tao_sinh_vien('SV040', self.lop_hanh_chinh)
        self.sinh_vien_khac = tao_sinh_vien('SV041', self.lop_hanh_chinh)
        self.giang_vien = tao_giang_vien('GV040')
        self.hoc_phan = HocPhan.objects.create(ma_hoc_phan='CTDL', ten_hoc_phan='Cấu trúc dữ liệu', so_tin_chi=3)
        self.lop_hoc_phan = LopHocPhan.objects.create(
            ma_lop_hoc_phan='CTDL_01', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=40,
            ma_hoc_phan=self.hoc_phan, ma_giang_vien=self.giang_vien,
        )
        self.dang_ky = DangKyHocPhan.objects.create(
            ma_sinh_vien=self.sinh_vien, ma_lop_hoc_phan=self.lop_hoc_phan, thoi_gian_dang_ky=timezone.now(),
        )

    def test_huy_thanh_cong_qua_service(self):
        huy_dang_ky_hoc_phan(self.dang_ky)
        self.dang_ky.refresh_from_db()
        self.assertEqual(self.dang_ky.trang_thai, TrangThaiDangKy.DA_HUY)
        # Dòng vẫn còn, không bị xóa (giữ lịch sử).
        self.assertTrue(DangKyHocPhan.objects.filter(pk=self.dang_ky.pk).exists())

    def test_khong_the_huy_hai_lan(self):
        huy_dang_ky_hoc_phan(self.dang_ky)
        with self.assertRaises(LoiDangKyHocPhan) as ngu_canh:
            huy_dang_ky_hoc_phan(self.dang_ky)
        self.assertIn('đã được hủy', str(ngu_canh.exception))

    def test_chan_huy_khi_da_co_diem(self):
        Diem.objects.create(ma_dang_ky=self.dang_ky, diem_qua_trinh=8, diem_thi=8)
        with self.assertRaises(LoiDangKyHocPhan) as ngu_canh:
            huy_dang_ky_hoc_phan(self.dang_ky)
        self.assertIn('đã có điểm', str(ngu_canh.exception))
        self.dang_ky.refresh_from_db()
        self.assertEqual(self.dang_ky.trang_thai, TrangThaiDangKy.CON_HIEU_LUC)

    def test_van_huy_duoc_khi_diem_rong_chua_nhap_gi(self):
        # nhap_diem_view tạo Diem rỗng (get_or_create) ngay khi giảng viên mở trang, dù
        # chưa nhập điểm nào — trường hợp này vẫn phải hủy được bình thường.
        Diem.objects.create(ma_dang_ky=self.dang_ky, diem_qua_trinh=None, diem_thi=None)
        huy_dang_ky_hoc_phan(self.dang_ky)
        self.dang_ky.refresh_from_db()
        self.assertEqual(self.dang_ky.trang_thai, TrangThaiDangKy.DA_HUY)

    def test_dang_ky_lai_sau_khi_huy_tao_dong_moi(self):
        huy_dang_ky_hoc_phan(self.dang_ky)
        dang_ky_moi = dang_ky_hoc_phan(self.sinh_vien, self.lop_hoc_phan)
        self.assertNotEqual(dang_ky_moi.pk, self.dang_ky.pk)
        self.assertEqual(dang_ky_moi.trang_thai, TrangThaiDangKy.CON_HIEU_LUC)
        # Dòng cũ đã hủy vẫn còn nguyên trong lịch sử.
        self.assertEqual(
            DangKyHocPhan.objects.filter(ma_sinh_vien=self.sinh_vien, ma_lop_hoc_phan=self.lop_hoc_phan).count(), 2,
        )

    def test_huy_qua_view_thanh_cong_va_xac_nhan_truoc(self):
        self.client.force_login(self.sinh_vien.ma_tai_khoan)
        response = self.client.get(reverse('enrollment:huy_dang_ky', args=[self.dang_ky.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertNotEqual(
            DangKyHocPhan.objects.get(pk=self.dang_ky.pk).trang_thai, TrangThaiDangKy.DA_HUY,
        )

        response = self.client.post(reverse('enrollment:huy_dang_ky', args=[self.dang_ky.pk]))
        self.assertRedirects(response, reverse('enrollment:lich_su_dang_ky'))
        self.dang_ky.refresh_from_db()
        self.assertEqual(self.dang_ky.trang_thai, TrangThaiDangKy.DA_HUY)

    def test_khong_the_huy_dang_ky_cua_nguoi_khac(self):
        dang_ky_cua_khac = DangKyHocPhan.objects.create(
            ma_sinh_vien=self.sinh_vien_khac, ma_lop_hoc_phan=self.lop_hoc_phan, thoi_gian_dang_ky=timezone.now(),
        )
        self.client.force_login(self.sinh_vien.ma_tai_khoan)
        response = self.client.post(reverse('enrollment:huy_dang_ky', args=[dang_ky_cua_khac.pk]))
        self.assertEqual(response.status_code, 403)
        dang_ky_cua_khac.refresh_from_db()
        self.assertEqual(dang_ky_cua_khac.trang_thai, TrangThaiDangKy.CON_HIEU_LUC)

    def test_giang_vien_khong_duoc_huy_dang_ky(self):
        self.client.force_login(self.giang_vien.ma_tai_khoan)
        response = self.client.post(reverse('enrollment:huy_dang_ky', args=[self.dang_ky.pk]))
        self.assertEqual(response.status_code, 403)


class SiSoVaThongKeSauKhiHuyTest(TestCase):
    """Phần B: sĩ số và các kiểm tra khác chỉ tính đăng ký còn hiệu lực sau khi hủy."""

    def setUp(self):
        tao_dot_dang_ky_mo('1', '2025-2026')
        self.lop_hanh_chinh = tao_lop_hanh_chinh('CNTT1')
        self.giang_vien = tao_giang_vien('GV050')
        self.hoc_phan = HocPhan.objects.create(ma_hoc_phan='CTDL', ten_hoc_phan='Cấu trúc dữ liệu', so_tin_chi=3)
        self.lop_hoc_phan = LopHocPhan.objects.create(
            ma_lop_hoc_phan='CTDL_01', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=1,
            ma_hoc_phan=self.hoc_phan, ma_giang_vien=self.giang_vien,
        )
        self.sv1 = tao_sinh_vien('SV050', self.lop_hanh_chinh)
        self.sv2 = tao_sinh_vien('SV051', self.lop_hanh_chinh)
        self.dang_ky_sv1 = DangKyHocPhan.objects.create(
            ma_sinh_vien=self.sv1, ma_lop_hoc_phan=self.lop_hoc_phan, thoi_gian_dang_ky=timezone.now(),
        )

    def test_si_so_giam_sau_khi_huy_cho_phep_sinh_vien_khac_dang_ky(self):
        # Lớp sĩ số tối đa 1, đã đầy bởi sv1 -> sv2 bị chặn.
        with self.assertRaises(LoiDangKyHocPhan) as ngu_canh:
            dang_ky_hoc_phan(self.sv2, self.lop_hoc_phan)
        self.assertIn('sĩ số', str(ngu_canh.exception))

        huy_dang_ky_hoc_phan(self.dang_ky_sv1)

        dang_ky_sv2 = dang_ky_hoc_phan(self.sv2, self.lop_hoc_phan)
        self.assertTrue(DangKyHocPhan.objects.filter(pk=dang_ky_sv2.pk).exists())

    def test_trang_danh_sach_lop_dang_mo_giam_si_so_sau_khi_huy(self):
        huy_dang_ky_hoc_phan(self.dang_ky_sv1)
        self.client.force_login(self.sv2.ma_tai_khoan)
        response = self.client.get(reverse('enrollment:lop_dang_mo'))
        # Trang Sinh viên hiện sĩ số dạng "đã đăng ký/tối đa" trong chip (báo cáo của Phòng Đào tạo vẫn "x / y").
        self.assertContains(response, '0/1')

    def test_bao_cao_si_so_theo_lop_chi_tinh_con_hieu_luc(self):
        pdt = TaiKhoan.objects.create_user(ten_dang_nhap='pdt_ss', vai_tro=VaiTro.PHONG_DAO_TAO, password=MAT_KHAU)
        huy_dang_ky_hoc_phan(self.dang_ky_sv1)
        self.client.force_login(pdt)
        response = self.client.get(reverse('reports:si_so_theo_lop'))
        self.assertContains(response, '0 / 1')

    def test_bao_cao_hoc_phan_nhieu_dang_ky_chi_tinh_con_hieu_luc(self):
        pdt = TaiKhoan.objects.create_user(ten_dang_nhap='pdt_hp', vai_tro=VaiTro.PHONG_DAO_TAO, password=MAT_KHAU)
        huy_dang_ky_hoc_phan(self.dang_ky_sv1)
        self.client.force_login(pdt)
        response = self.client.get(reverse('reports:hoc_phan_nhieu_dang_ky'))
        self.assertEqual(response.status_code, 200)
        hoc_phan = next(hp for hp in response.context['danh_sach'] if hp.pk == self.hoc_phan.pk)
        self.assertEqual(hoc_phan.so_dang_ky, 0)

    def test_ket_qua_hoc_tap_bo_qua_dang_ky_da_huy(self):
        Diem.objects.create(ma_dang_ky=self.dang_ky_sv1, diem_qua_trinh=8, diem_thi=8)
        # sv2 đăng ký (thẳng qua ORM, bỏ qua kiểm tra sĩ số) rồi hủy TRƯỚC KHI có điểm —
        # đăng ký đã hủy này không được tính vào báo cáo dù sau đó có tạo điểm cho nó.
        dang_ky_sv2 = DangKyHocPhan.objects.create(
            ma_sinh_vien=self.sv2, ma_lop_hoc_phan=self.lop_hoc_phan, thoi_gian_dang_ky=timezone.now(),
        )
        huy_dang_ky_hoc_phan(dang_ky_sv2)
        # Bình thường không thể xảy ra qua giao diện (nhap_diem_view chặn chấm điểm cho
        # đăng ký đã hủy), nhưng tạo thẳng qua ORM ở đây để xác nhận rõ ràng report LỌC
        # theo trang_thai chứ không chỉ tình cờ đúng vì thiếu điểm.
        Diem.objects.create(ma_dang_ky=dang_ky_sv2, diem_qua_trinh=9, diem_thi=9)

        pdt = TaiKhoan.objects.create_user(ten_dang_nhap='pdt_kq', vai_tro=VaiTro.PHONG_DAO_TAO, password=MAT_KHAU)
        self.client.force_login(pdt)
        response = self.client.get(reverse('reports:ket_qua_hoc_tap'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Phân bố điểm tổng kết (1 lượt có đủ điểm)')


class TrungLichTienQuyetBoQuaDaHuyTest(TestCase):
    """Phần B: kiểm tra trùng lịch/tiên quyết khi đăng ký học phần khác chỉ tính các đăng
    ký còn hiệu lực — một đăng ký đã hủy không được coi là "đã đạt" hay "đã chiếm lịch"."""

    def setUp(self):
        tao_dot_dang_ky_mo('1', '2025-2026')
        self.lop_hanh_chinh = tao_lop_hanh_chinh('CNTT1')
        self.sinh_vien = tao_sinh_vien('SV060', self.lop_hanh_chinh)
        self.giang_vien = tao_giang_vien('GV060')
        self.ctdl = HocPhan.objects.create(ma_hoc_phan='CTDL', ten_hoc_phan='Cấu trúc dữ liệu', so_tin_chi=3)
        self.csdl = HocPhan.objects.create(ma_hoc_phan='CSDL', ten_hoc_phan='Cơ sở dữ liệu', so_tin_chi=3)
        DieuKienTienQuyet.objects.create(ma_hoc_phan=self.csdl, ma_hoc_phan_tien_quyet=self.ctdl)
        self.lop_csdl = LopHocPhan.objects.create(
            ma_lop_hoc_phan='CSDL_01', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=40,
            ma_hoc_phan=self.csdl, ma_giang_vien=self.giang_vien,
        )

    def test_huy_dang_ky_lop_trung_lich_roi_dang_ky_lop_moi_khong_bao_trung(self):
        # csdl yêu cầu tiên quyết ctdl (setUp) — phải cho sinh viên đạt ctdl trước, để
        # phép thử chỉ còn vướng đúng trùng lịch (kiểm tra tiên quyết chạy trước trùng lịch).
        lop_ctdl_cu = LopHocPhan.objects.create(
            ma_lop_hoc_phan='CTDL_CU', hoc_ky='1', nam_hoc='2024-2025', si_so_toi_da=40,
            ma_hoc_phan=self.ctdl, ma_giang_vien=self.giang_vien,
        )
        dang_ky_ctdl_cu = DangKyHocPhan.objects.create(
            ma_sinh_vien=self.sinh_vien, ma_lop_hoc_phan=lop_ctdl_cu, thoi_gian_dang_ky=timezone.now(),
        )
        Diem.objects.create(ma_dang_ky=dang_ky_ctdl_cu, diem_qua_trinh=8, diem_thi=8)

        lop_khac = LopHocPhan.objects.create(
            ma_lop_hoc_phan='KHAC_01', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=40,
            ma_hoc_phan=self.ctdl, ma_giang_vien=self.giang_vien,
        )
        LichHoc.objects.create(ma_lop_hoc_phan=self.lop_csdl, thu=2, tiet_bat_dau=1, so_tiet=3, phong_hoc='A101')
        LichHoc.objects.create(ma_lop_hoc_phan=lop_khac, thu=2, tiet_bat_dau=1, so_tiet=3, phong_hoc='B101')
        dang_ky_khac = DangKyHocPhan.objects.create(
            ma_sinh_vien=self.sinh_vien, ma_lop_hoc_phan=lop_khac, thoi_gian_dang_ky=timezone.now(),
        )

        with self.assertRaises(LoiDangKyHocPhan) as ngu_canh:
            dang_ky_hoc_phan(self.sinh_vien, self.lop_csdl)
        self.assertIn('trùng', str(ngu_canh.exception))

        huy_dang_ky_hoc_phan(dang_ky_khac)
        dang_ky_moi = dang_ky_hoc_phan(self.sinh_vien, self.lop_csdl)
        self.assertTrue(DangKyHocPhan.objects.filter(pk=dang_ky_moi.pk).exists())

    def test_dang_ky_da_huy_du_co_diem_khong_tinh_la_da_dat_tien_quyet(self):
        # Bình thường không xảy ra qua giao diện (hủy bị chặn nếu đã có điểm), nhưng đặt
        # thẳng trang_thai=DA_HUY qua ORM ở đây để xác nhận rõ tiên quyết CHỈ tính đăng ký
        # còn hiệu lực, không phải "chỉ tình cờ đúng vì đăng ký đã hủy luôn không có điểm".
        lop_ctdl_cu = LopHocPhan.objects.create(
            ma_lop_hoc_phan='CTDL_CU', hoc_ky='1', nam_hoc='2024-2025', si_so_toi_da=40,
            ma_hoc_phan=self.ctdl, ma_giang_vien=self.giang_vien,
        )
        dang_ky_cu = DangKyHocPhan.objects.create(
            ma_sinh_vien=self.sinh_vien, ma_lop_hoc_phan=lop_ctdl_cu, thoi_gian_dang_ky=timezone.now(),
            trang_thai=TrangThaiDangKy.DA_HUY,
        )
        Diem.objects.create(ma_dang_ky=dang_ky_cu, diem_qua_trinh=8, diem_thi=8)

        with self.assertRaises(LoiDangKyHocPhan) as ngu_canh:
            dang_ky_hoc_phan(self.sinh_vien, self.lop_csdl)
        self.assertIn('tiên quyết', str(ngu_canh.exception))


class DanhSachSinhVienLopTrangThaiTest(TestCase):
    """Phòng Đào tạo phải thấy đủ trạng thái (không ẩn đăng ký đã hủy); giảng viên chỉ
    thấy sinh viên đang thực học (đăng ký còn hiệu lực)."""

    def setUp(self):
        self.lop_hanh_chinh = tao_lop_hanh_chinh('CNTT1')
        self.giang_vien = tao_giang_vien('GV070')
        self.phong_dao_tao = TaiKhoan.objects.create_user(
            ten_dang_nhap='pdt070', vai_tro=VaiTro.PHONG_DAO_TAO, password=MAT_KHAU
        )
        self.hoc_phan = HocPhan.objects.create(ma_hoc_phan='CTDL', ten_hoc_phan='Cấu trúc dữ liệu', so_tin_chi=3)
        self.lop_hoc_phan = LopHocPhan.objects.create(
            ma_lop_hoc_phan='CTDL_01', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=40,
            ma_hoc_phan=self.hoc_phan, ma_giang_vien=self.giang_vien,
        )
        self.sv_con_hieu_luc = tao_sinh_vien('SV070', self.lop_hanh_chinh)
        self.sv_da_huy = tao_sinh_vien('SV071', self.lop_hanh_chinh)
        DangKyHocPhan.objects.create(
            ma_sinh_vien=self.sv_con_hieu_luc, ma_lop_hoc_phan=self.lop_hoc_phan, thoi_gian_dang_ky=timezone.now(),
        )
        dang_ky_da_huy = DangKyHocPhan.objects.create(
            ma_sinh_vien=self.sv_da_huy, ma_lop_hoc_phan=self.lop_hoc_phan, thoi_gian_dang_ky=timezone.now(),
        )
        huy_dang_ky_hoc_phan(dang_ky_da_huy)

    def test_phong_dao_tao_thay_ca_da_huy_va_trang_thai(self):
        self.client.force_login(self.phong_dao_tao)
        response = self.client.get(reverse('enrollment:danh_sach_sinh_vien_lop', args=[self.lop_hoc_phan.pk]))
        self.assertContains(response, 'SV070')
        self.assertContains(response, 'SV071')
        self.assertContains(response, 'Còn hiệu lực')
        self.assertContains(response, 'Đã hủy')

    def test_giang_vien_chi_thay_con_hieu_luc(self):
        self.client.force_login(self.giang_vien.ma_tai_khoan)
        response = self.client.get(reverse('enrollment:danh_sach_sinh_vien_lop', args=[self.lop_hoc_phan.pk]))
        self.assertContains(response, 'SV070')
        self.assertNotContains(response, 'SV071')

    def test_giang_vien_xuat_excel_chi_gom_con_hieu_luc(self):
        self.client.force_login(self.giang_vien.ma_tai_khoan)
        response = self.client.get(reverse('enrollment:xuat_excel_sinh_vien_lop', args=[self.lop_hoc_phan.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        sheet = load_workbook(BytesIO(response.content)).active
        ma_cac_sinh_vien = [row[0].value for row in sheet.iter_rows(min_row=2)]
        self.assertEqual(ma_cac_sinh_vien, ['SV070'])

    def test_phong_dao_tao_xuat_excel_gom_ca_da_huy(self):
        self.client.force_login(self.phong_dao_tao)
        response = self.client.get(reverse('enrollment:xuat_excel_sinh_vien_lop', args=[self.lop_hoc_phan.pk]))
        sheet = load_workbook(BytesIO(response.content)).active
        ma_cac_sinh_vien = {row[0].value for row in sheet.iter_rows(min_row=2)}
        self.assertEqual(ma_cac_sinh_vien, {'SV070', 'SV071'})

    def test_giang_vien_khac_khong_xuat_excel_duoc(self):
        giang_vien_khac = tao_giang_vien('GV071')
        self.client.force_login(giang_vien_khac.ma_tai_khoan)
        response = self.client.get(reverse('enrollment:xuat_excel_sinh_vien_lop', args=[self.lop_hoc_phan.pk]))
        self.assertEqual(response.status_code, 403)


class XuatExcelSinhVienPhuTrachTest(TestCase):
    def setUp(self):
        self.lop_hanh_chinh = tao_lop_hanh_chinh('CNTT090')
        self.sinh_vien = tao_sinh_vien('SV090X', self.lop_hanh_chinh)
        self.co_van = self.lop_hanh_chinh.ma_co_van
        self.client.force_login(self.co_van.ma_tai_khoan)

    def test_xuat_excel_danh_sach_sinh_vien_phu_trach(self):
        response = self.client.get(reverse('enrollment:xuat_excel_sinh_vien_phu_trach'))
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        ma_cac_sinh_vien = [row[0].value for row in sheet.iter_rows(min_row=2)]
        self.assertEqual(ma_cac_sinh_vien, ['SV090X'])

    def test_sinh_vien_khong_xuat_excel_duoc(self):
        self.client.force_login(self.sinh_vien.ma_tai_khoan)
        response = self.client.get(reverse('enrollment:xuat_excel_sinh_vien_phu_trach'))
        self.assertEqual(response.status_code, 403)


class LuoiTuanSinhVienTest(TestCase):
    def setUp(self):
        tao_dot_dang_ky_mo('1', '2025-2026')
        self.lop_hanh_chinh = tao_lop_hanh_chinh('CNTT1')
        self.sinh_vien = tao_sinh_vien('SV080', self.lop_hanh_chinh)
        self.giang_vien = tao_giang_vien('GV080')
        hoc_phan = HocPhan.objects.create(ma_hoc_phan='CSDL', ten_hoc_phan='Cơ sở dữ liệu', so_tin_chi=3)
        self.lop = LopHocPhan.objects.create(
            ma_lop_hoc_phan='CSDL_LUOI_SV', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=40,
            ma_hoc_phan=hoc_phan, ma_giang_vien=self.giang_vien,
        )
        LichHoc.objects.create(ma_lop_hoc_phan=self.lop, thu=3, tiet_bat_dau=2, so_tiet=2, phong_hoc='B202')
        dang_ky_hoc_phan(self.sinh_vien, self.lop)
        self.client.force_login(self.sinh_vien.ma_tai_khoan)

    def test_xem_dang_luoi_tuan_hien_dung_lop_va_phong(self):
        response = self.client.get(reverse('enrollment:thoi_khoa_bieu'), {'xem': 'luoi'})
        self.assertContains(response, 'CSDL_LUOI_SV')
        self.assertContains(response, 'B202')

    def test_mac_dinh_van_la_dang_danh_sach(self):
        response = self.client.get(reverse('enrollment:thoi_khoa_bieu'))
        self.assertNotContains(response, '<table class="bang-luoi-tuan">')

    def test_xuat_pdf_chay_khong_loi(self):
        response = self.client.get(reverse('enrollment:thoi_khoa_bieu_xuat_pdf'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        noi_dung = b''.join(response.streaming_content)
        self.assertTrue(noi_dung.startswith(b'%PDF'))


class LichThiSinhVienTest(TestCase):
    def setUp(self):
        tao_dot_dang_ky_mo('1', '2025-2026')
        self.lop_hanh_chinh = tao_lop_hanh_chinh('CNTT1')
        self.sinh_vien = tao_sinh_vien('SV090', self.lop_hanh_chinh)
        self.giang_vien = tao_giang_vien('GV090')
        hoc_phan = HocPhan.objects.create(ma_hoc_phan='CSDL', ten_hoc_phan='Cơ sở dữ liệu', so_tin_chi=3)
        hoc_phan_khac = HocPhan.objects.create(ma_hoc_phan='CTDL', ten_hoc_phan='Cấu trúc dữ liệu', so_tin_chi=3)

        self.lop_dang_hoc = LopHocPhan.objects.create(
            ma_lop_hoc_phan='CSDL_THI01', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=40,
            ma_hoc_phan=hoc_phan, ma_giang_vien=self.giang_vien,
        )
        self.lop_da_huy = LopHocPhan.objects.create(
            ma_lop_hoc_phan='CTDL_THI01', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=40,
            ma_hoc_phan=hoc_phan_khac, ma_giang_vien=self.giang_vien,
        )
        LichThi.objects.create(
            lop_hoc_phan=self.lop_dang_hoc, ngay_thi=date(2026, 1, 20), gio_bat_dau=time(8, 0),
            thoi_luong_phut=60, phong_thi='A301',
        )
        LichThi.objects.create(
            lop_hoc_phan=self.lop_da_huy, ngay_thi=date(2026, 1, 22), gio_bat_dau=time(9, 0),
            thoi_luong_phut=60, phong_thi='A302',
        )

        dang_ky_hoc_phan(self.sinh_vien, self.lop_dang_hoc)
        dang_ky_da_huy = dang_ky_hoc_phan(self.sinh_vien, self.lop_da_huy)
        huy_dang_ky_hoc_phan(dang_ky_da_huy)

        self.client.force_login(self.sinh_vien.ma_tai_khoan)

    def test_chi_thay_lich_thi_cua_lop_con_hieu_luc(self):
        response = self.client.get(reverse('enrollment:lich_thi'))
        self.assertContains(response, 'CSDL_THI01')
        self.assertContains(response, 'A301')
        self.assertNotContains(response, 'CTDL_THI01')
        self.assertNotContains(response, 'A302')

    def test_xuat_pdf_chay_khong_loi(self):
        response = self.client.get(reverse('enrollment:lich_thi_xuat_pdf'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        noi_dung = b''.join(response.streaming_content)
        self.assertTrue(noi_dung.startswith(b'%PDF'))

    def test_xuat_pdf_khong_co_lich_thi_van_chay_khong_loi(self):
        sinh_vien_khac = tao_sinh_vien('SV091', self.lop_hanh_chinh)
        self.client.force_login(sinh_vien_khac.ma_tai_khoan)
        response = self.client.get(reverse('enrollment:lich_thi_xuat_pdf'))
        self.assertEqual(response.status_code, 200)
        noi_dung = b''.join(response.streaming_content)
        self.assertTrue(noi_dung.startswith(b'%PDF'))


def _the_a_trong_main(response):
    phan_main = re.search(r'<main>.*?</main>', response.content.decode(), re.S).group(0)
    return re.findall(r'<a [^>]*>', phan_main)


class LopDangMoGiaoDienTest(TestCase):
    """Giao diện trang "Lớp đang mở" của Sinh viên (lượt giao diện 3/3): sĩ số dạng chip theo
    mức đầy, chip trạng thái đầy đủ (tiên quyết, trùng lịch...), nút, dòng nhóm học kỳ."""

    def setUp(self):
        tao_dot_dang_ky_mo('1', '2025-2026')
        self.lop_hanh_chinh = tao_lop_hanh_chinh('CNTT1')
        self.sinh_vien = tao_sinh_vien('SV300', self.lop_hanh_chinh)
        self.nguoi_khac = [tao_sinh_vien(f'SV3{so:02d}', self.lop_hanh_chinh) for so in range(1, 6)]
        self.giang_vien = tao_giang_vien('GV300')
        self.hp = {
            ma: HocPhan.objects.create(ma_hoc_phan=ma, ten_hoc_phan=f'Học phần {ma}', so_tin_chi=3)
            for ma in ('HPA', 'HPB', 'HPC', 'HPD', 'HPE', 'HPF', 'HPG')
        }
        # HPB cần đạt HPA làm tiên quyết; sinh viên chưa đạt HPA.
        DieuKienTienQuyet.objects.create(ma_hoc_phan=self.hp['HPB'], ma_hoc_phan_tien_quyet=self.hp['HPA'])
        self.lop = {}
        for ma_lop, hoc_phan, si_so in (
            ('L_TRONG', 'HPC', 10), ('L_SAP_DAY', 'HPD', 5), ('L_DAY', 'HPE', 1),
            ('L_TIEN_QUYET', 'HPB', 10), ('L_TRUNG', 'HPF', 10), ('L_DA_DK', 'HPG', 10),
        ):
            self.lop[ma_lop] = LopHocPhan.objects.create(
                ma_lop_hoc_phan=ma_lop, hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=si_so,
                ma_hoc_phan=self.hp[hoc_phan], ma_giang_vien=self.giang_vien,
            )
        for sinh_vien in self.nguoi_khac[:4]:
            self._dang_ky(sinh_vien, 'L_SAP_DAY')
        self._dang_ky(self.nguoi_khac[4], 'L_DAY')
        self._dang_ky(self.sinh_vien, 'L_DA_DK')
        LichHoc.objects.create(ma_lop_hoc_phan=self.lop['L_DA_DK'], thu=2, tiet_bat_dau=1, so_tiet=3, phong_hoc='A1')
        LichHoc.objects.create(ma_lop_hoc_phan=self.lop['L_TRUNG'], thu=2, tiet_bat_dau=2, so_tiet=2, phong_hoc='A2')
        self.client.force_login(self.sinh_vien.ma_tai_khoan)
        self.url = reverse('enrollment:lop_dang_mo')

    def _dang_ky(self, sinh_vien, ma_lop):
        return DangKyHocPhan.objects.create(
            ma_sinh_vien=sinh_vien, ma_lop_hoc_phan=self.lop[ma_lop], thoi_gian_dang_ky=timezone.now(),
        )

    def _trang(self, **tham_so):
        return self.client.get(self.url, tham_so)

    def test_si_so_dang_chip_theo_muc_day(self):
        noi_dung = self._trang().content.decode()
        self.assertIn('<span class="chip chip--du_dieu_kien" title="Còn nhiều chỗ">0/10</span>', noi_dung)
        self.assertIn('<span class="chip chip--thieu_tien_quyet" title="Sắp đầy">4/5</span>', noi_dung)
        self.assertIn('<span class="chip chip--da_day" title="Đã đầy">1/1</span>', noi_dung)

    def test_moi_lop_co_dung_trang_thai(self):
        trang_thai = {lop.pk: lop.trang_thai_dk['trang_thai'] for lop in self._trang().context['danh_sach']}
        self.assertEqual(trang_thai, {
            'L_TRONG': 'du_dieu_kien', 'L_SAP_DAY': 'du_dieu_kien', 'L_DAY': 'da_day',
            'L_TIEN_QUYET': 'thieu_tien_quyet', 'L_TRUNG': 'trung_lich', 'L_DA_DK': 'da_dang_ky',
        })

    def test_chip_trang_thai_dung_kieu_co_san(self):
        response = self._trang()
        for chip in (
            '<span class="chip chip--da_dang_ky">Đã đăng ký</span>',
            '<span class="chip chip--thieu_tien_quyet">Thiếu tiên quyết</span>',
            '<span class="chip chip--trung_lich">Trùng lịch</span>',
            '<span class="chip chip--da_day">Đã đầy</span>',
            '<span class="chip chip--du_dieu_kien">Đủ điều kiện</span>',
        ):
            self.assertContains(response, chip)
        self.assertContains(response, 'Cần đạt: Học phần HPA')
        self.assertContains(response, 'Với lớp L_DA_DK (Thứ 2, tiết 1, 3 tiết)')

    def test_chi_lop_du_dieu_kien_co_nut_dang_ky_chinh_con_lai_la_nut_vo_hieu(self):
        response = self._trang()
        for ma_lop in ('L_TRONG', 'L_SAP_DAY'):
            self.assertContains(response, f'action="{reverse("enrollment:dang_ky", args=[ma_lop])}"')
        for ma_lop in ('L_DAY', 'L_TIEN_QUYET', 'L_TRUNG', 'L_DA_DK'):
            self.assertNotContains(response, f'action="{reverse("enrollment:dang_ky", args=[ma_lop])}"')
        self.assertContains(response, '<button type="submit" class="nut nut-chinh nut-nho">Đăng ký</button>', count=2)
        self.assertContains(response, '<button type="button" class="nut nut-vo-hieu nut-nho" disabled>Đăng ký</button>', count=4)

    def test_dong_nhom_hoc_ky_neu_ro_dot_dang_mo(self):
        response = self._trang()
        self.assertContains(response, '<tr class="hang-nhom">', count=1)
        self.assertContains(response, 'Học kỳ 1 - 2025-2026')
        self.assertContains(response, '<span class="chip chip--du_dieu_kien">Đang mở đăng ký</span>')

    def test_ly_do_ngoai_dot_chi_nem_mot_lan_o_dong_nhom(self):
        # Trước đây câu lý do hiện ở từng hàng (và trùng với thông báo lỗi sau khi bấm đăng ký).
        DotDangKy.objects.all().delete()
        lop_them = LopHocPhan.objects.create(
            ma_lop_hoc_phan='L_THEM', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=10,
            ma_hoc_phan=self.hp['HPC'], ma_giang_vien=self.giang_vien,
        )
        response = self._trang()
        ly_do = 'Chưa có đợt đăng ký nào cho học kỳ này.'
        self.assertEqual(len(response.context['danh_sach']), 7)
        self.assertContains(response, ly_do, count=1)
        self.assertContains(response, '<span class="chip chip--ngoai_dot_dang_ky">Ngoài đợt đăng ký</span>')
        self.assertNotContains(response, f'action="{reverse("enrollment:dang_ky", args=[lop_them.pk])}"')

    def test_moi_nhom_hoc_ky_co_dong_nhom_rieng(self):
        HocKy.objects.create(
            hoc_ky=2, nam_hoc='2025-2026', ngay_bat_dau=timezone.localdate(),
            ngay_ket_thuc=timezone.localdate() + timedelta(days=100),
        )
        LopHocPhan.objects.create(
            ma_lop_hoc_phan='L_KY2', hoc_ky='2', nam_hoc='2025-2026', si_so_toi_da=10,
            ma_hoc_phan=self.hp['HPC'], ma_giang_vien=self.giang_vien,
        )
        response = self._trang()
        self.assertContains(response, '<tr class="hang-nhom">', count=2)
        self.assertContains(response, 'Học kỳ 2 - 2025-2026')
        self.assertContains(response, 'Chưa có đợt đăng ký nào cho học kỳ này.', count=1)

    def test_tai_khoan_chua_gan_ho_so_khong_co_nut_dang_ky(self):
        tk = TaiKhoan.objects.create_user(ten_dang_nhap='sv_khong_ho_so', vai_tro=VaiTro.SINH_VIEN, password=MAT_KHAU)
        self.client.force_login(tk)
        response = self._trang()
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Tài khoản của bạn chưa gắn hồ sơ sinh viên.')
        self.assertNotContains(response, '<button type="submit" class="nut nut-chinh nut-nho">Đăng ký</button>')

    def test_bo_loc_giu_nguyen_logic(self):
        response = self._trang(con_cho='1')
        ma_cac_lop = {lop.pk for lop in response.context['danh_sach']}
        self.assertNotIn('L_DAY', ma_cac_lop)
        self.assertIn('L_TRONG', ma_cac_lop)

    def test_moi_hanh_dong_la_nut_khong_gach_chan_va_html_can_bang(self):
        response = self._trang()
        for the in _the_a_trong_main(response):
            self.assertIn('class="nut', the)
        self.assertNotIn('underline', response.content.decode())
        loi, con_mo = kiem_tra_can_bang_the(response.content.decode())
        self.assertEqual((loi, con_mo), ([], []))


class ThoiKhoaBieuVaLichSuGiaoDienTest(TestCase):
    def setUp(self):
        tao_dot_dang_ky_mo('1', '2025-2026')
        self.lop_hanh_chinh = tao_lop_hanh_chinh('CNTT1')
        self.sinh_vien = tao_sinh_vien('SV310', self.lop_hanh_chinh)
        self.giang_vien = tao_giang_vien('GV310')
        hoc_phan = HocPhan.objects.create(ma_hoc_phan='CSDL', ten_hoc_phan='Cơ sở dữ liệu', so_tin_chi=3)
        hoc_phan_huy = HocPhan.objects.create(ma_hoc_phan='CTDL', ten_hoc_phan='Cấu trúc dữ liệu', so_tin_chi=3)
        self.lop = LopHocPhan.objects.create(
            ma_lop_hoc_phan='CSDL_TKB', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=40,
            ma_hoc_phan=hoc_phan, ma_giang_vien=self.giang_vien,
        )
        self.lop_huy = LopHocPhan.objects.create(
            ma_lop_hoc_phan='CTDL_TKB', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=40,
            ma_hoc_phan=hoc_phan_huy, ma_giang_vien=self.giang_vien,
        )
        LichHoc.objects.create(ma_lop_hoc_phan=self.lop, thu=3, tiet_bat_dau=2, so_tiet=2, phong_hoc='B202')
        dang_ky_hoc_phan(self.sinh_vien, self.lop)
        dang_ky_huy = dang_ky_hoc_phan(self.sinh_vien, self.lop_huy)
        huy_dang_ky_hoc_phan(dang_ky_huy)
        self.client.force_login(self.sinh_vien.ma_tai_khoan)

    def test_thoi_khoa_bieu_nut_xuat_pdf_chinh_canh_tieu_de_va_nut_huy_nguy_hiem(self):
        response = self.client.get(reverse('enrollment:thoi_khoa_bieu'))
        phan_dau = re.search(r'<div class="tieu-de-trang">.*?</div>\s*</div>', response.content.decode(), re.S).group(0)
        self.assertIn(f'<a class="nut nut-chinh" href="{reverse("enrollment:thoi_khoa_bieu_xuat_pdf")}?ky=1%3A2025-2026">Xuất PDF</a>', phan_dau)
        self.assertContains(response, '<a class="nut nut-phu" href="?xem=luoi&amp;ky=1%3A2025-2026">Xem dạng lưới tuần</a>')
        self.assertContains(response, 'class="nut nut-nguy-hiem nut-nho"')
        self.assertContains(response, '<table class="bang-sv">')
        for the in _the_a_trong_main(response):
            self.assertIn('class="nut', the)

    def test_luoi_tuan_o_co_lop_dung_lop_va_o_trong_khong_co_lop(self):
        response = self.client.get(reverse('enrollment:thoi_khoa_bieu'), {'xem': 'luoi'})
        noi_dung = response.content.decode()
        self.assertContains(response, '<table class="bang-luoi-tuan bang-sv">')
        # Thứ 3, tiết 2-3: hai ô có lớp (mỗi tiết một ô); các ô khác đều trống.
        self.assertEqual(noi_dung.count('<div class="o-lich-hoc">'), 2)
        self.assertContains(response, '<a class="nut nut-phu" href="?ky=1%3A2025-2026">Xem dạng danh sách</a>')
        css = re.search(r'\.o-lich-hoc \{.*?\}', noi_dung, re.S).group(0)
        self.assertIn('background: var(--mau-xanh-nhat);', css)
        self.assertIn('border-left: 3px solid var(--mau-xanh-chinh);', css)

    def test_lich_su_dang_ky_chip_va_lam_nhat_dong_da_huy(self):
        response = self.client.get(reverse('enrollment:lich_su_dang_ky'))
        noi_dung = response.content.decode()
        self.assertContains(response, '<span class="chip chip--con_hieu_luc">Còn hiệu lực</span>')
        self.assertContains(response, '<span class="chip chip--da_huy">Đã hủy</span>')
        self.assertEqual(noi_dung.count('<tr class="hang-diem hang-mo">'), 1)
        self.assertEqual(noi_dung.count('<tr class="hang-diem">'), 1)
        dong_da_huy = re.search(r'<tr class="hang-diem hang-mo">.*?</tr>', noi_dung, re.S).group(0)
        self.assertIn('CTDL_TKB', dong_da_huy)
        self.assertNotIn('Hủy đăng ký', dong_da_huy)
        self.assertContains(response, '<table class="bang-sv">')


class LichThiGiaoDienTest(TestCase):
    def setUp(self):
        tao_dot_dang_ky_mo('1', '2025-2026')
        self.hom_nay = timezone.localdate()
        self.lop_hanh_chinh = tao_lop_hanh_chinh('CNTT1')
        self.sinh_vien = tao_sinh_vien('SV320', self.lop_hanh_chinh)
        self.giang_vien = tao_giang_vien('GV320')
        self.client.force_login(self.sinh_vien.ma_tai_khoan)

    def _tao_lich_thi(self, ma, ngay, gio='08:00'):
        hoc_phan = HocPhan.objects.create(ma_hoc_phan=ma, ten_hoc_phan=f'Học phần {ma}', so_tin_chi=3)
        lop = LopHocPhan.objects.create(
            ma_lop_hoc_phan=f'{ma}_L', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=40,
            ma_hoc_phan=hoc_phan, ma_giang_vien=self.giang_vien,
        )
        DangKyHocPhan.objects.create(
            ma_sinh_vien=self.sinh_vien, ma_lop_hoc_phan=lop, thoi_gian_dang_ky=timezone.now(),
        )
        return LichThi.objects.create(
            lop_hoc_phan=lop, ngay_thi=ngay, gio_bat_dau=gio, thoi_luong_phut=90, phong_thi=f'P_{ma}',
        )

    def test_gom_theo_ngay_thi_va_lam_nhat_lich_da_qua(self):
        ngay_qua, ngay_1, ngay_2 = (self.hom_nay + timedelta(days=d) for d in (-3, 5, 9))
        self._tao_lich_thi('QUA', ngay_qua)
        self._tao_lich_thi('AAA', ngay_1, '08:00')
        self._tao_lich_thi('BBB', ngay_1, '13:30')
        self._tao_lich_thi('CCC', ngay_2)
        response = self.client.get(reverse('enrollment:lich_thi'))
        noi_dung = response.content.decode()
        self.assertEqual(noi_dung.count('<tr class="hang-nhom">'), 3)
        # Dòng nhóm có dạng "Thứ 5, 15/10/2026": hai lịch cùng ngày chỉ có MỘT dòng nhóm.
        self.assertContains(response, f', {ngay_1:%d/%m/%Y}</td></tr>', count=1)
        # Thứ tự: lịch đã qua, rồi hai lịch cùng ngày (theo giờ), rồi ngày sau.
        self.assertEqual(
            [lich.lop_hoc_phan_id for lich in response.context['danh_sach']], ['QUA_L', 'AAA_L', 'BBB_L', 'CCC_L'],
        )
        self.assertEqual(noi_dung.count('<tr class="hang-diem hang-mo">'), 1)
        dong_qua = re.search(r'<tr class="hang-diem hang-mo">.*?</tr>', noi_dung, re.S).group(0)
        self.assertIn('QUA_L', dong_qua)

    def test_cac_cot_chi_gom_du_lieu_that_khong_co_sbd(self):
        self._tao_lich_thi('AAA', self.hom_nay + timedelta(days=2))
        response = self.client.get(reverse('enrollment:lich_thi'))
        phan_dau = re.search(r'<thead>.*?</thead>', response.content.decode(), re.S).group(0)
        cac_cot = [re.sub(r'<[^>]+>', '', cot).strip() for cot in re.findall(r'<th[^>]*>.*?</th>', phan_dau, re.S)]
        self.assertEqual(cac_cot, ['Lớp học phần', 'Học phần', 'Giờ bắt đầu', 'Thời lượng (phút)', 'Phòng thi'])
        self.assertNotContains(response, 'SBD')
        self.assertContains(response, '<td class="trung-tam">08:00</td>')
        self.assertContains(response, '<td>P_AAA</td>')

    def test_nut_xuat_pdf_la_nut_chinh_va_khong_co_lien_ket_chu(self):
        response = self.client.get(reverse('enrollment:lich_thi'))
        self.assertContains(response, f'<a class="nut nut-chinh" href="{reverse("enrollment:lich_thi_xuat_pdf")}">Xuất PDF</a>')
        for the in _the_a_trong_main(response):
            self.assertIn('class="nut', the)
        self.assertContains(response, 'Bạn chưa có lịch thi nào.')
        loi, con_mo = kiem_tra_can_bang_the(response.content.decode())
        self.assertEqual((loi, con_mo), ([], []))


class VuaCotPdfTest(TestCase):
    """Chữ trong bảng PDF không được tràn sang cột bên cạnh (lỗi đè chữ ở thời khóa biểu)."""

    def setUp(self):
        from reportlab.lib.units import cm
        from reportlab.pdfbase import pdfmetrics

        from .pdf_utils import dam_bao_font_dejavu, vua_cot
        dam_bao_font_dejavu()
        self.cm, self.do_rong_chu, self.vua_cot = cm, pdfmetrics.stringWidth, vua_cot

    def test_chu_ngan_giu_nguyen_co_chu(self):
        chu, co = self.vua_cot('Thứ 2', 'DejaVuSans', 10, 3 * self.cm)
        self.assertEqual((chu, co), ('Thứ 2', 10))

    def test_tieu_de_dai_thu_nho_cho_vua_cot(self):
        do_rong = 1.8 * self.cm
        chu, co = self.vua_cot('Tiết bắt đầu', 'DejaVuSans-Bold', 10, do_rong)
        self.assertEqual(chu, 'Tiết bắt đầu')
        self.assertLess(co, 10)
        self.assertLessEqual(self.do_rong_chu(chu, 'DejaVuSans-Bold', co), do_rong)

    def test_chu_qua_dai_bi_cat_kem_dau_ba_cham(self):
        do_rong = 2 * self.cm
        chu, co = self.vua_cot('Cấu trúc dữ liệu và giải thuật nâng cao', 'DejaVuSans', 10, do_rong)
        self.assertTrue(chu.endswith('…'))
        self.assertEqual(co, 7)
        self.assertLessEqual(self.do_rong_chu(chu, 'DejaVuSans', co), do_rong)

    def test_tieu_de_thoi_khoa_bieu_vua_cot_o_co_chu_goc(self):
        from .pdf_utils import _KHOANG_CACH_COT
        from .views import _CAC_COT_THOI_KHOA_BIEU
        for ten_cot, do_rong in _CAC_COT_THOI_KHOA_BIEU:
            _, co = self.vua_cot(ten_cot, 'DejaVuSans-Bold', 10, do_rong * self.cm - _KHOANG_CACH_COT)
            self.assertEqual(co, 10, ten_cot)


class ThoiKhoaBieuTheoHocKyTest(TestCase):
    """Thời khóa biểu (web và PDF) chỉ hiện lớp của một học kỳ: mặc định học kỳ đang diễn ra,
    chọn được học kỳ khác bằng ?ky=<hoc_ky>:<nam_hoc>."""

    def setUp(self):
        tao_dot_dang_ky_mo('2', '2025-2026')  # học kỳ đang diễn ra hôm nay
        HocKy.objects.create(
            hoc_ky=1, nam_hoc='2024-2025', ngay_bat_dau=date(2024, 9, 1), ngay_ket_thuc=date(2025, 1, 15),
        )
        self.sinh_vien = tao_sinh_vien('SV420', tao_lop_hanh_chinh('CNTT1'))
        giang_vien = tao_giang_vien('GV420')
        hoc_phan_moi = HocPhan.objects.create(ma_hoc_phan='MMT', ten_hoc_phan='Mạng máy tính', so_tin_chi=3)
        hoc_phan_cu = HocPhan.objects.create(ma_hoc_phan='TOAN', ten_hoc_phan='Toán cao cấp', so_tin_chi=3)
        self.lop_moi = LopHocPhan.objects.create(
            ma_lop_hoc_phan='MMT_KY_MOI', hoc_ky='2', nam_hoc='2025-2026', si_so_toi_da=40,
            ma_hoc_phan=hoc_phan_moi, ma_giang_vien=giang_vien,
        )
        self.lop_cu = LopHocPhan.objects.create(
            ma_lop_hoc_phan='TOAN_KY_CU', hoc_ky='1', nam_hoc='2024-2025', si_so_toi_da=40,
            ma_hoc_phan=hoc_phan_cu, ma_giang_vien=giang_vien,
        )
        LichHoc.objects.create(ma_lop_hoc_phan=self.lop_moi, thu=2, tiet_bat_dau=4, so_tiet=3, phong_hoc='C201')
        LichHoc.objects.create(ma_lop_hoc_phan=self.lop_cu, thu=2, tiet_bat_dau=1, so_tiet=3, phong_hoc='C101')
        DangKyHocPhan.objects.create(
            ma_sinh_vien=self.sinh_vien, ma_lop_hoc_phan=self.lop_cu, thoi_gian_dang_ky=timezone.now() - timedelta(days=400),
        )
        self.dang_ky_moi = DangKyHocPhan.objects.create(
            ma_sinh_vien=self.sinh_vien, ma_lop_hoc_phan=self.lop_moi, thoi_gian_dang_ky=timezone.now(),
        )
        self.client.force_login(self.sinh_vien.ma_tai_khoan)

    def _ma_lop(self, response):
        return [lich.ma_lop_hoc_phan_id for lich in response.context['danh_sach']]

    def test_mac_dinh_chi_hien_hoc_ky_dang_dien_ra(self):
        response = self.client.get(reverse('enrollment:thoi_khoa_bieu'))
        self.assertEqual(self._ma_lop(response), ['MMT_KY_MOI'])
        self.assertEqual(response.context['ky_chon'], '2:2025-2026')
        self.assertContains(response, '<option value="1:2024-2025" >Học kỳ 1 - 2024-2025</option>', html=False)

    def test_chon_hoc_ky_cu_chi_hien_lop_ky_cu(self):
        response = self.client.get(reverse('enrollment:thoi_khoa_bieu'), {'ky': '1:2024-2025'})
        self.assertEqual(self._ma_lop(response), ['TOAN_KY_CU'])

    def test_tham_so_ky_khong_hop_le_quay_ve_mac_dinh(self):
        response = self.client.get(reverse('enrollment:thoi_khoa_bieu'), {'ky': '9:1999-2000'})
        self.assertEqual(self._ma_lop(response), ['MMT_KY_MOI'])

    def test_khong_co_lop_ky_hien_tai_thi_lay_ky_gan_nhat_co_lop(self):
        self.dang_ky_moi.delete()
        response = self.client.get(reverse('enrollment:thoi_khoa_bieu'))
        self.assertEqual(self._ma_lop(response), ['TOAN_KY_CU'])

    def test_xuat_pdf_dung_hoc_ky_dang_chon(self):
        from unittest import mock

        from . import views
        with mock.patch.object(views, 'tao_pdf_bang', wraps=views.tao_pdf_bang) as gia:
            response = self.client.get(reverse('enrollment:thoi_khoa_bieu_xuat_pdf'), {'ky': '1:2024-2025'})
        self.assertEqual(response.status_code, 200)
        tieu_de, _, cac_dong = gia.call_args.args[:3]
        self.assertEqual(tieu_de, 'THỜI KHÓA BIỂU - HỌC KỲ 1 NĂM HỌC 2024-2025')
        self.assertEqual([dong[1] for dong in cac_dong], ['TOAN_KY_CU'])
