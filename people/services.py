from datetime import date, datetime

from django.contrib.auth import get_user_model
from django.db import IntegrityError, transaction
from openpyxl import load_workbook

from academics.models import LopHanhChinh
from accounts.models import VaiTro
from accounts.forms import ten_dang_nhap_hop_le, thong_bao_ten_dang_nhap_khong_hop_le
from accounts.services import tao_mat_khau_tam

from .lien_he import chuan_hoa_so_dien_thoai, kiem_tra_lien_he
from .ma_ho_so import ma_hop_le, thong_bao_ma_khong_hop_le
from .models import SinhVien

TaiKhoan = get_user_model()

TIEU_DE_COT_MAU_SINH_VIEN = [
    'Mã sinh viên', 'Họ tên', 'Ngày sinh (YYYY-MM-DD)', 'Email', 'Số điện thoại',
    'Mã lớp hành chính', 'Tên đăng nhập (để trống sẽ tự tạo từ mã sinh viên)',
]


def _doc_gia_tri(dong, chi_muc):
    if chi_muc >= len(dong):
        return None
    gia_tri = dong[chi_muc]
    if isinstance(gia_tri, str):
        gia_tri = gia_tri.strip()
    return gia_tri or None


def _chuan_hoa_ngay_sinh(gia_tri):
    if isinstance(gia_tri, datetime):
        return gia_tri.date()
    if isinstance(gia_tri, date):
        return gia_tri
    if isinstance(gia_tri, str):
        for dinh_dang in ('%Y-%m-%d', '%d/%m/%Y'):
            try:
                return datetime.strptime(gia_tri, dinh_dang).date()
            except ValueError:
                continue
    return None


def xu_ly_file_excel_sinh_vien(file_excel):
    """Đọc file Excel danh sách sinh viên (đúng thứ tự cột TIEU_DE_COT_MAU_SINH_VIEN),
    kiểm tra từng dòng, chỉ lưu các dòng hợp lệ — dòng lỗi bị bỏ qua kèm lý do theo số dòng.

    Mỗi dòng hợp lệ được tạo kèm một tài khoản đăng nhập mới (mật khẩu tạm ngẫu nhiên,
    bắt buộc đổi ngay lần đăng nhập đầu — giống accounts.dat_lai_mat_khau), vì cột
    ma_tai_khoan của sinh_vien không cho NULL. Tên đăng nhập lấy từ cột cuối nếu có, để
    trống thì tự tạo bằng mã sinh viên viết thường.

    Trả về dict {'so_thanh_cong': int, 'loi': [{'dong', 'loi'}],
    'tao_moi': [{'ma_sinh_vien', 'ten_dang_nhap', 'mat_khau_tam'}]}. mat_khau_tam chỉ để
    hiển thị đúng một lần ở trang kết quả của request này, không lưu lại ở đâu khác.
    """
    workbook = load_workbook(file_excel, data_only=True)
    sheet = workbook.active

    loi = []
    tao_moi = []
    ma_sinh_vien_da_dung = set(SinhVien.objects.values_list('ma_sinh_vien', flat=True))
    ten_dang_nhap_da_dung = set(TaiKhoan.objects.values_list('ten_dang_nhap', flat=True))
    # Email (chữ thường) / số điện thoại (đã chuẩn hóa) của các dòng hợp lệ trước đó trong tệp -> số dòng.
    email_trong_tep, sdt_trong_tep = {}, {}

    for so_dong, dong in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if dong is None or all(gia_tri in (None, '') for gia_tri in dong):
            continue

        ma_sinh_vien = _doc_gia_tri(dong, 0)
        ho_ten = _doc_gia_tri(dong, 1)
        ngay_sinh_tho = _doc_gia_tri(dong, 2)
        email = _doc_gia_tri(dong, 3)
        so_dien_thoai = _doc_gia_tri(dong, 4)
        ma_lop = _doc_gia_tri(dong, 5)
        ten_dang_nhap = _doc_gia_tri(dong, 6)

        thieu = [
            ten for ten, gia_tri in [
                ('mã sinh viên', ma_sinh_vien), ('họ tên', ho_ten), ('ngày sinh', ngay_sinh_tho),
                ('email', email), ('số điện thoại', so_dien_thoai), ('mã lớp hành chính', ma_lop),
            ] if not gia_tri
        ]
        if thieu:
            loi.append({'dong': so_dong, 'loi': f'Thiếu cột bắt buộc: {", ".join(thieu)}.'})
            continue

        if isinstance(ma_sinh_vien, float) and ma_sinh_vien.is_integer():
            ma_sinh_vien = int(ma_sinh_vien)  # ô số trong Excel (2023001.0) -> 2023001
        ma_sinh_vien = str(ma_sinh_vien).strip()
        if not ma_hop_le(ma_sinh_vien):
            loi.append({'dong': so_dong, 'loi': thong_bao_ma_khong_hop_le('Mã sinh viên', ma_sinh_vien)})
            continue
        if ma_sinh_vien in ma_sinh_vien_da_dung:
            loi.append({'dong': so_dong, 'loi': f'Mã sinh viên "{ma_sinh_vien}" đã tồn tại.'})
            continue

        ngay_sinh = _chuan_hoa_ngay_sinh(ngay_sinh_tho)
        if ngay_sinh is None:
            loi.append({'dong': so_dong, 'loi': 'Ngày sinh không hợp lệ (dùng định dạng YYYY-MM-DD).'})
            continue

        ma_lop = str(ma_lop).strip()
        if not LopHanhChinh.objects.filter(pk=ma_lop).exists():
            loi.append({'dong': so_dong, 'loi': f'Mã lớp hành chính "{ma_lop}" không tồn tại.'})
            continue

        ten_dang_nhap = str(ten_dang_nhap).strip() if ten_dang_nhap else ma_sinh_vien.lower()
        if not ten_dang_nhap_hop_le(ten_dang_nhap):
            loi.append({'dong': so_dong, 'loi': thong_bao_ten_dang_nhap_khong_hop_le(ten_dang_nhap)})
            continue
        if ten_dang_nhap in ten_dang_nhap_da_dung:
            loi.append({'dong': so_dong, 'loi': f'Tên đăng nhập "{ten_dang_nhap}" đã tồn tại.'})
            continue

        email, so_dien_thoai = str(email).strip(), str(so_dien_thoai).strip()
        khoa_email, khoa_sdt = email.lower(), chuan_hoa_so_dien_thoai(so_dien_thoai)
        loi_lien_he = list(kiem_tra_lien_he(email, so_dien_thoai).values())
        if khoa_email in email_trong_tep:
            loi_lien_he.append(f'Email "{email}" trùng với dòng {email_trong_tep[khoa_email]} trong tệp.')
        if khoa_sdt in sdt_trong_tep:
            loi_lien_he.append(f'Số điện thoại "{so_dien_thoai}" trùng với dòng {sdt_trong_tep[khoa_sdt]} trong tệp.')
        if loi_lien_he:
            loi.append({'dong': so_dong, 'loi': ' '.join(loi_lien_he)})
            continue

        try:
            with transaction.atomic():
                tai_khoan_tam = TaiKhoan(ten_dang_nhap=ten_dang_nhap, vai_tro=VaiTro.SINH_VIEN)
                mat_khau_tam = tao_mat_khau_tam(tai_khoan_tam)
                tai_khoan = TaiKhoan.objects.create_user(
                    ten_dang_nhap=ten_dang_nhap, vai_tro=VaiTro.SINH_VIEN, password=mat_khau_tam,
                )
                tai_khoan.phai_doi_mat_khau = True
                tai_khoan.save(update_fields=['phai_doi_mat_khau'])
                SinhVien.objects.create(
                    ma_sinh_vien=ma_sinh_vien, ho_ten=str(ho_ten).strip(), ngay_sinh=ngay_sinh,
                    email=email, so_dien_thoai=so_dien_thoai,
                    ma_lop_id=ma_lop, ma_tai_khoan=tai_khoan,
                )
        except IntegrityError as loi_ghi:
            loi.append({'dong': so_dong, 'loi': f'Lỗi khi lưu vào cơ sở dữ liệu: {loi_ghi}'})
            continue

        ma_sinh_vien_da_dung.add(ma_sinh_vien)
        ten_dang_nhap_da_dung.add(ten_dang_nhap)
        email_trong_tep[khoa_email] = so_dong
        sdt_trong_tep[khoa_sdt] = so_dong
        tao_moi.append({'ma_sinh_vien': ma_sinh_vien, 'ten_dang_nhap': ten_dang_nhap, 'mat_khau_tam': mat_khau_tam})

    return {'so_thanh_cong': len(tao_moi), 'loi': loi, 'tao_moi': tao_moi}
