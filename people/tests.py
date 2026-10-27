from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, tag
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook

from academics.models import HocPhan, LopHanhChinh, NganhHoc
from accounts.models import VaiTro
from classes.models import LopHocPhan
from diem_danh.models import DiemDanh
from enrollment.models import DangKyHocPhan, TrangThaiDangKy

from .models import CoVanHocTap, GiangVien, SinhVien

TaiKhoan = get_user_model()
MAT_KHAU = 'MatKhauManh123'


def _tao_lop_hanh_chinh(ma_lop='CNTT1'):
    nganh, _ = NganhHoc.objects.get_or_create(ma_nganh='CNTT', defaults={'ten_nganh': 'Công nghệ thông tin'})
    tk_co_van = TaiKhoan.objects.create_user(
        ten_dang_nhap=f'cv_{ma_lop}', vai_tro=VaiTro.CO_VAN_HOC_TAP, password=MAT_KHAU
    )
    co_van = CoVanHocTap.objects.create(
        ma_co_van=f'CV_{ma_lop}', ho_ten='Cố vấn Test', email='cv@test.local',
        so_dien_thoai='0900000000', ma_tai_khoan=tk_co_van,
    )
    return LopHanhChinh.objects.create(ma_lop=ma_lop, ten_lop=f'Lớp {ma_lop}', ma_nganh=nganh, ma_co_van=co_van)


@tag('tich_hop')
class SinhVienCrudTest(TestCase):
    def setUp(self):
        self.phong_dao_tao = TaiKhoan.objects.create_user(
            ten_dang_nhap='pdt', vai_tro=VaiTro.PHONG_DAO_TAO, password=MAT_KHAU
        )
        self.sinh_vien_tk = TaiKhoan.objects.create_user(
            ten_dang_nhap='sv001', vai_tro=VaiTro.SINH_VIEN, password=MAT_KHAU
        )
        self.lop = _tao_lop_hanh_chinh('CNTT1')
        self.client.force_login(self.phong_dao_tao)

    def test_tao_sinh_vien(self):
        response = self.client.post(reverse('people:sinh_vien_tao'), {
            'ma_sinh_vien': 'SV001',
            'ho_ten': 'Nguyễn Văn A',
            'ngay_sinh': '2004-01-01',
            'email': 'a@test.local',
            'so_dien_thoai': '0900000001',
            'ma_lop': self.lop.pk,
            'ma_tai_khoan': self.sinh_vien_tk.pk,
        })
        self.assertRedirects(response, reverse('people:sinh_vien_danh_sach'))
        self.assertTrue(SinhVien.objects.filter(pk='SV001').exists())

    def test_khong_the_gan_tai_khoan_da_dung_cho_sinh_vien_khac(self):
        SinhVien.objects.create(
            ma_sinh_vien='SV002', ho_ten='B', ngay_sinh='2004-01-01', email='b@test.local',
            so_dien_thoai='0900000002', ma_lop=self.lop, ma_tai_khoan=self.sinh_vien_tk,
        )
        response = self.client.post(reverse('people:sinh_vien_tao'), {
            'ma_sinh_vien': 'SV003',
            'ho_ten': 'C',
            'ngay_sinh': '2004-01-01',
            'email': 'c@test.local',
            'so_dien_thoai': '0900000003',
            'ma_lop': self.lop.pk,
            'ma_tai_khoan': self.sinh_vien_tk.pk,
        })
        self.assertEqual(response.status_code, 200)
        self.assertFalse(SinhVien.objects.filter(pk='SV003').exists())

    def test_khong_the_chon_tai_khoan_sai_vai_tro(self):
        response = self.client.post(reverse('people:sinh_vien_tao'), {
            'ma_sinh_vien': 'SV004',
            'ho_ten': 'D',
            'ngay_sinh': '2004-01-01',
            'email': 'd@test.local',
            'so_dien_thoai': '0900000004',
            'ma_lop': self.lop.pk,
            'ma_tai_khoan': self.phong_dao_tao.pk,
        })
        self.assertEqual(response.status_code, 200)
        self.assertFalse(SinhVien.objects.filter(pk='SV004').exists())

    def test_tim_kiem_theo_ma_va_ho_ten(self):
        SinhVien.objects.create(
            ma_sinh_vien='SV005', ho_ten='Trần Thị E', ngay_sinh='2004-01-01', email='e@test.local',
            so_dien_thoai='0900000005', ma_lop=self.lop, ma_tai_khoan=self.sinh_vien_tk,
        )
        response = self.client.get(reverse('people:sinh_vien_danh_sach'), {'q': 'Thị E'})
        self.assertContains(response, 'SV005')

    def test_loc_theo_lop_hanh_chinh(self):
        lop_khac = _tao_lop_hanh_chinh('CNTT2')
        SinhVien.objects.create(
            ma_sinh_vien='SV006', ho_ten='F', ngay_sinh='2004-01-01', email='f@test.local',
            so_dien_thoai='0900000006', ma_lop=lop_khac, ma_tai_khoan=self.sinh_vien_tk,
        )
        response = self.client.get(reverse('people:sinh_vien_danh_sach'), {'ma_lop': self.lop.pk})
        self.assertNotContains(response, 'SV006')


def _tao_file_excel(cac_dong):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['Mã sinh viên', 'Họ tên', 'Ngày sinh', 'Email', 'Số điện thoại', 'Mã lớp hành chính', 'Tên đăng nhập'])
    for dong in cac_dong:
        sheet.append(dong)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return SimpleUploadedFile(
        'sinh_vien.xlsx', buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@tag('tich_hop')
class XoaHoSoTest(TestCase):
    """Xóa hồ sơ sinh viên/giảng viên/cố vấn: chỉ khi chưa có dữ liệu đi kèm; tài khoản đăng nhập vẫn còn."""

    def setUp(self):
        self.phong_dao_tao = TaiKhoan.objects.create_user(ten_dang_nhap='pdt_xhs', vai_tro=VaiTro.PHONG_DAO_TAO, password=MAT_KHAU)
        self.client.force_login(self.phong_dao_tao)
        self.lop = _tao_lop_hanh_chinh('XHS')  # đã có cố vấn CV_XHS đang phụ trách lớp này

        self.tk_sv = TaiKhoan.objects.create_user(ten_dang_nhap='sv_xhs', vai_tro=VaiTro.SINH_VIEN, password=MAT_KHAU)
        self.sinh_vien = SinhVien.objects.create(
            ma_sinh_vien='SV_XHS', ho_ten='Sinh viên XHS', ngay_sinh='2004-01-01', email='sv@test.local',
            so_dien_thoai='0900000000', ma_lop=self.lop, ma_tai_khoan=self.tk_sv,
        )
        self.tk_gv = TaiKhoan.objects.create_user(ten_dang_nhap='gv_xhs', vai_tro=VaiTro.GIANG_VIEN, password=MAT_KHAU)
        self.giang_vien = GiangVien.objects.create(
            ma_giang_vien='GV_XHS', ho_ten='Giảng viên XHS', email='gv@test.local', so_dien_thoai='0900000000',
            ma_tai_khoan=self.tk_gv,
        )
        self.tk_cv = TaiKhoan.objects.create_user(ten_dang_nhap='cv_tu_do', vai_tro=VaiTro.CO_VAN_HOC_TAP, password=MAT_KHAU)
        self.co_van_tu_do = CoVanHocTap.objects.create(
            ma_co_van='CV_TU_DO', ho_ten='Cố vấn chưa có lớp', email='cv@test.local', so_dien_thoai='0900000000',
            ma_tai_khoan=self.tk_cv,
        )
        self.url_sv = reverse('people:sinh_vien_xoa', args=['SV_XHS'])
        self.url_gv = reverse('people:giang_vien_xoa', args=['GV_XHS'])
        self.url_cv = reverse('people:co_van_xoa', args=['CV_TU_DO'])

    def _lop_hoc_phan(self):
        hoc_phan = HocPhan.objects.create(ma_hoc_phan='HP_XHS', ten_hoc_phan='HP', so_tin_chi=3)
        return LopHocPhan.objects.create(
            ma_lop_hoc_phan='LHP_XHS', hoc_ky='1', nam_hoc='2025-2026', si_so_toi_da=10,
            ma_hoc_phan=hoc_phan, ma_giang_vien=self.giang_vien,
        )

    def test_cac_danh_sach_co_nut_xoa(self):
        for ten_url, url in (
            ('people:sinh_vien_danh_sach', self.url_sv), ('people:giang_vien_danh_sach', self.url_gv),
            ('people:co_van_danh_sach', self.url_cv),
        ):
            self.assertContains(self.client.get(reverse(ten_url)), f'href="{url}"', msg_prefix=ten_url)

    def test_trang_xac_nhan_khong_xoa_va_neu_ro_tai_khoan_con_lai(self):
        for url, ma, tai_khoan in ((self.url_sv, 'SV_XHS', 'sv_xhs'), (self.url_gv, 'GV_XHS', 'gv_xhs'), (self.url_cv, 'CV_TU_DO', 'cv_tu_do')):
            response = self.client.get(url)
            self.assertContains(response, ma)
            self.assertContains(response, f'tài khoản đăng nhập "{tai_khoan}" vẫn còn')
            self.assertNotContains(response, 'chưa thể xóa')
        self.assertTrue(SinhVien.objects.filter(pk='SV_XHS').exists())

    def test_xoa_sinh_vien_chua_dung_giu_lai_tai_khoan(self):
        response = self.client.post(self.url_sv, follow=True)
        self.assertRedirects(response, reverse('people:sinh_vien_danh_sach'))
        self.assertContains(response, 'Đã xóa hồ sơ sinh viên')
        self.assertFalse(SinhVien.objects.filter(pk='SV_XHS').exists())
        self.assertTrue(TaiKhoan.objects.filter(ten_dang_nhap='sv_xhs').exists())

    def test_chan_xoa_sinh_vien_co_dang_ky_ke_ca_da_huy(self):
        lop = self._lop_hoc_phan()
        DangKyHocPhan.objects.create(
            ma_sinh_vien=self.sinh_vien, ma_lop_hoc_phan=lop, thoi_gian_dang_ky=timezone.now(), trang_thai=TrangThaiDangKy.DA_HUY,
        )
        response = self.client.post(self.url_sv, follow=True)
        self.assertRedirects(response, reverse('people:sinh_vien_danh_sach'))
        self.assertContains(response, '1 bản đăng ký học phần')
        self.assertContains(response, 'khóa tài khoản')
        self.assertTrue(SinhVien.objects.filter(pk='SV_XHS').exists())
        self.assertContains(self.client.get(self.url_sv), 'chưa thể xóa')

    def test_chan_xoa_sinh_vien_co_diem_danh(self):
        lop = self._lop_hoc_phan()
        DiemDanh.objects.create(lop_hoc_phan=lop, sinh_vien=self.sinh_vien, ngay_diem_danh=timezone.localdate())
        response = self.client.post(self.url_sv, follow=True)
        self.assertContains(response, '1 bản ghi điểm danh')
        self.assertTrue(SinhVien.objects.filter(pk='SV_XHS').exists())

    def test_xoa_giang_vien_chua_dung_va_chan_khi_co_lop_hoc_phan(self):
        self._lop_hoc_phan()
        response = self.client.post(self.url_gv, follow=True)
        self.assertContains(response, '1 lớp học phần đang phụ trách')
        self.assertTrue(GiangVien.objects.filter(pk='GV_XHS').exists())
        LopHocPhan.objects.all().delete()
        response = self.client.post(self.url_gv, follow=True)
        self.assertRedirects(response, reverse('people:giang_vien_danh_sach'))
        self.assertContains(response, 'Đã xóa hồ sơ giảng viên')
        self.assertFalse(GiangVien.objects.filter(pk='GV_XHS').exists())
        self.assertTrue(TaiKhoan.objects.filter(ten_dang_nhap='gv_xhs').exists())

    def test_xoa_co_van_chua_dung_va_chan_khi_dang_phu_trach_lop_hanh_chinh(self):
        LopHanhChinh.objects.filter(pk=self.lop.pk).update(ma_co_van=self.co_van_tu_do)
        response = self.client.post(self.url_cv, follow=True)
        self.assertContains(response, '1 lớp hành chính đang phụ trách')
        self.assertTrue(CoVanHocTap.objects.filter(pk='CV_TU_DO').exists())
        LopHanhChinh.objects.filter(pk=self.lop.pk).update(ma_co_van=CoVanHocTap.objects.get(pk='CV_XHS'))
        response = self.client.post(self.url_cv, follow=True)
        self.assertRedirects(response, reverse('people:co_van_danh_sach'))
        self.assertFalse(CoVanHocTap.objects.filter(pk='CV_TU_DO').exists())
        self.assertTrue(TaiKhoan.objects.filter(ten_dang_nhap='cv_tu_do').exists())

    def test_sau_khi_xoa_tai_khoan_khong_gan_ho_so_van_gan_lai_duoc(self):
        self.client.post(self.url_sv)
        response = self.client.post(reverse('people:sinh_vien_tao'), {
            'ma_sinh_vien': 'SVMOI', 'ho_ten': 'Sinh viên mới', 'ngay_sinh': '2004-02-02', 'email': 'moi@test.local',
            'so_dien_thoai': '0900000001', 'ma_lop': self.lop.pk, 'ma_tai_khoan': self.tk_sv.pk,
        })
        self.assertRedirects(response, reverse('people:sinh_vien_danh_sach'))
        self.assertTrue(SinhVien.objects.filter(pk='SVMOI', ma_tai_khoan=self.tk_sv).exists())

    def test_ho_so_khong_ton_tai_tra_ve_404(self):
        for ten_url in ('people:sinh_vien_xoa', 'people:giang_vien_xoa', 'people:co_van_xoa'):
            self.assertEqual(self.client.post(reverse(ten_url, args=['KHONG_CO'])).status_code, 404, ten_url)

    def test_vai_tro_khac_va_chua_dang_nhap_khong_xoa_duoc(self):
        cac_url = (self.url_sv, self.url_gv, self.url_cv)
        for vai_tro in (VaiTro.GIANG_VIEN, VaiTro.SINH_VIEN, VaiTro.CO_VAN_HOC_TAP):
            self.client.force_login(TaiKhoan.objects.create_user(ten_dang_nhap=f'q_{vai_tro}', vai_tro=vai_tro, password=MAT_KHAU))
            for url in cac_url:
                self.assertEqual(self.client.post(url).status_code, 403, (vai_tro, url))
        self.client.logout()
        for url in cac_url:
            self.assertEqual(self.client.post(url).status_code, 302, url)
        self.assertTrue(SinhVien.objects.filter(pk='SV_XHS').exists())
        self.assertTrue(GiangVien.objects.filter(pk='GV_XHS').exists())
        self.assertTrue(CoVanHocTap.objects.filter(pk='CV_TU_DO').exists())


@tag('tich_hop')
class NhapSinhVienExcelTest(TestCase):
    def setUp(self):
        self.phong_dao_tao = TaiKhoan.objects.create_user(
            ten_dang_nhap='pdt_excel', vai_tro=VaiTro.PHONG_DAO_TAO, password=MAT_KHAU
        )
        self.lop = _tao_lop_hanh_chinh('CNTT_EXCEL1')
        self.client.force_login(self.phong_dao_tao)

    def test_nhap_du_lieu_hop_le_tao_ca_tai_khoan(self):
        file_excel = _tao_file_excel([
            ['SVE001', 'Nguyễn Văn X', '2004-05-01', 'x@test.local', '0911111111', self.lop.pk, ''],
        ])
        response = self.client.post(reverse('people:sinh_vien_nhap_excel'), {'file_excel': file_excel})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['ket_qua']['so_thanh_cong'], 1)
        self.assertEqual(response.context['ket_qua']['loi'], [])

        sinh_vien = SinhVien.objects.get(pk='SVE001')
        self.assertEqual(sinh_vien.ma_tai_khoan.ten_dang_nhap, 'sve001')
        self.assertTrue(sinh_vien.ma_tai_khoan.phai_doi_mat_khau)

    def test_dong_thieu_cot_bat_buoc_bi_bao_loi(self):
        file_excel = _tao_file_excel([
            ['SVE002', '', '2004-05-01', 'y@test.local', '0911111112', self.lop.pk, ''],
        ])
        response = self.client.post(reverse('people:sinh_vien_nhap_excel'), {'file_excel': file_excel})
        self.assertFalse(SinhVien.objects.filter(pk='SVE002').exists())
        self.assertEqual(len(response.context['ket_qua']['loi']), 1)
        self.assertIn('Thiếu cột bắt buộc', response.context['ket_qua']['loi'][0]['loi'])
        self.assertEqual(response.context['ket_qua']['loi'][0]['dong'], 2)

    def test_dong_ma_lop_khong_ton_tai_bi_bao_loi(self):
        file_excel = _tao_file_excel([
            ['SVE003', 'Z', '2004-05-01', 'z@test.local', '0911111113', 'LOP_KHONG_CO', ''],
        ])
        response = self.client.post(reverse('people:sinh_vien_nhap_excel'), {'file_excel': file_excel})
        self.assertFalse(SinhVien.objects.filter(pk='SVE003').exists())
        self.assertIn('không tồn tại', response.context['ket_qua']['loi'][0]['loi'])

    def test_dong_ma_sinh_vien_trung_bi_loi_dong_khac_van_thanh_cong(self):
        SinhVien.objects.create(
            ma_sinh_vien='SVE004', ho_ten='Đã có', ngay_sinh='2004-01-01', email='old@test.local',
            so_dien_thoai='0900000000', ma_lop=self.lop,
            ma_tai_khoan=TaiKhoan.objects.create_user(
                ten_dang_nhap='sve004_cu', vai_tro=VaiTro.SINH_VIEN, password=MAT_KHAU,
            ),
        )
        file_excel = _tao_file_excel([
            ['SVE004', 'Trùng', '2004-05-01', 'trung@test.local', '0911111114', self.lop.pk, ''],
            ['SVE005', 'Hợp lệ', '2004-05-01', 'hople@test.local', '0911111115', self.lop.pk, ''],
        ])
        response = self.client.post(reverse('people:sinh_vien_nhap_excel'), {'file_excel': file_excel})
        self.assertEqual(response.context['ket_qua']['so_thanh_cong'], 1)
        self.assertIn('đã tồn tại', response.context['ket_qua']['loi'][0]['loi'])
        self.assertTrue(SinhVien.objects.filter(pk='SVE005').exists())

    def test_dong_ten_dang_nhap_chi_dinh_duoc_su_dung(self):
        file_excel = _tao_file_excel([
            ['SVE006', 'Có tên đăng nhập', '2004-05-01', 'w@test.local', '0911111116', self.lop.pk, 'nguoidung_w'],
        ])
        self.client.post(reverse('people:sinh_vien_nhap_excel'), {'file_excel': file_excel})
        sinh_vien = SinhVien.objects.get(pk='SVE006')
        self.assertEqual(sinh_vien.ma_tai_khoan.ten_dang_nhap, 'nguoidung_w')

    def test_tai_mau_excel(self):
        response = self.client.get(reverse('people:sinh_vien_mau_excel'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

    def test_sinh_vien_khong_vao_duoc_trang_nhap_excel(self):
        sinh_vien = TaiKhoan.objects.create_user(ten_dang_nhap='sv_excel', vai_tro=VaiTro.SINH_VIEN, password=MAT_KHAU)
        self.client.force_login(sinh_vien)
        response = self.client.get(reverse('people:sinh_vien_nhap_excel'))
        self.assertEqual(response.status_code, 403)


@tag('tich_hop')
class PhanQuyenNguoiDungTest(TestCase):
    def setUp(self):
        self.sinh_vien = TaiKhoan.objects.create_user(
            ten_dang_nhap='sv999', vai_tro=VaiTro.SINH_VIEN, password=MAT_KHAU
        )

    def test_chua_dang_nhap_bi_chuyen_huong(self):
        response = self.client.get(reverse('people:sinh_vien_danh_sach'))
        self.assertEqual(response.status_code, 302)

    def test_sinh_vien_khong_vao_duoc_danh_sach_sinh_vien(self):
        self.client.force_login(self.sinh_vien)
        response = self.client.get(reverse('people:sinh_vien_danh_sach'))
        self.assertEqual(response.status_code, 403)


MA_KHONG_HOP_LE = ['///', ']]]', 'SV 01', 'SV-01', 'SV_01', 'Nguyễn01', 'a/b', '']


@tag('don_vi')
class QuyTacMaHoSoTest(TestCase):
    def test_chi_chap_nhan_chu_cai_khong_dau_va_chu_so(self):
        from .ma_ho_so import ma_hop_le
        for ma in ('SV001', 'sv001', '2023001', 'GV1', 'A'):
            self.assertTrue(ma_hop_le(ma), ma)
        for ma in MA_KHONG_HOP_LE + ['SV001\n', None, 123]:
            self.assertFalse(ma_hop_le(ma), repr(ma))

    def test_dung_duoc_trong_url_chi_loai_ma_lam_hong_lien_ket(self):
        from .ma_ho_so import ma_dung_duoc_trong_url
        # Dữ liệu cũ có "_", "]" hay dấu cách vẫn tạo liên kết được (Django mã hóa %xx); "/", ".", ".." thì không.
        for ma in ('SV001', 'SV_01', ']]]', 'SV 01'):
            self.assertTrue(ma_dung_duoc_trong_url(ma), ma)
        for ma in ('///', 'a/b', '', '.', '..', None):
            self.assertFalse(ma_dung_duoc_trong_url(ma), repr(ma))


@tag('tich_hop')
class MaHoSoFormTest(TestCase):
    def setUp(self):
        self.phong_dao_tao = TaiKhoan.objects.create_user(
            ten_dang_nhap='pdt_ma', vai_tro=VaiTro.PHONG_DAO_TAO, password=MAT_KHAU
        )
        self.lop = _tao_lop_hanh_chinh('CNTTMA')
        self.client.force_login(self.phong_dao_tao)

    def _tai_khoan(self, ten, vai_tro):
        return TaiKhoan.objects.create_user(ten_dang_nhap=ten, vai_tro=vai_tro, password=MAT_KHAU)

    def test_form_sinh_vien_tu_choi_ma_co_ky_tu_la(self):
        tk = self._tai_khoan('sv_ma', VaiTro.SINH_VIEN)
        for ma in MA_KHONG_HOP_LE:
            response = self.client.post(reverse('people:sinh_vien_tao'), {
                'ma_sinh_vien': ma, 'ho_ten': 'A', 'ngay_sinh': '2004-01-01', 'email': 'ma@test.local',
                'so_dien_thoai': '0912000001', 'ma_lop': self.lop.pk, 'ma_tai_khoan': tk.pk,
            })
            self.assertEqual(response.status_code, 200, ma)
            self.assertFalse(SinhVien.objects.filter(ma_tai_khoan=tk).exists(), ma)
            if ma:
                self.assertContains(response, 'chỉ được dùng chữ cái không dấu và chữ số')

    def test_form_giang_vien_va_co_van_tu_choi_ma_co_ky_tu_la_nhan_ma_hop_le(self):
        tk_gv = self._tai_khoan('gv_ma', VaiTro.GIANG_VIEN)
        tk_cv = self._tai_khoan('cv_ma', VaiTro.CO_VAN_HOC_TAP)
        du_lieu = {'ho_ten': 'B', 'email': 'gvma@test.local', 'so_dien_thoai': '0912000002'}
        self.client.post(reverse('people:giang_vien_tao'), {**du_lieu, 'ma_giang_vien': '///', 'ma_tai_khoan': tk_gv.pk})
        self.client.post(reverse('people:co_van_tao'), {**du_lieu, 'ma_co_van': ']]]', 'ma_tai_khoan': tk_cv.pk})
        self.assertFalse(GiangVien.objects.filter(ma_tai_khoan=tk_gv).exists())
        self.assertFalse(CoVanHocTap.objects.filter(ma_tai_khoan=tk_cv).exists())
        response = self.client.post(reverse('people:giang_vien_tao'), {**du_lieu, 'ma_giang_vien': 'GV01', 'ma_tai_khoan': tk_gv.pk})
        self.assertRedirects(response, reverse('people:giang_vien_danh_sach'))

    def test_sua_ho_so_cu_co_ma_khong_hop_le_van_sua_duoc_truong_khac(self):
        tk = self._tai_khoan('sv_cu', VaiTro.SINH_VIEN)
        SinhVien.objects.create(
            ma_sinh_vien='SV_CU', ho_ten='Cũ', ngay_sinh='2004-01-01', email='cu@test.local',
            so_dien_thoai='0912000003', ma_lop=self.lop, ma_tai_khoan=tk,
        )
        from .forms import SinhVienForm
        form = SinhVienForm({
            'ma_sinh_vien': 'SV_CU', 'ho_ten': 'Cũ đã sửa', 'ngay_sinh': '2004-01-01', 'email': 'cu@test.local',
            'so_dien_thoai': '0912000003', 'ma_lop': self.lop.pk, 'ma_tai_khoan': tk.pk,
        }, instance=SinhVien.objects.get(pk='SV_CU'))
        self.assertTrue(form.is_valid(), form.errors)


@tag('tich_hop')
class DuLieuCuMaKhongHopLeKhongSapTrangTest(TestCase):
    """Dữ liệu cũ (tạo trước khi có ràng buộc) có mã chứa "/" từng làm {% url %} ném NoReverseMatch."""

    def setUp(self):
        self.phong_dao_tao = TaiKhoan.objects.create_user(
            ten_dang_nhap='pdt_cu', vai_tro=VaiTro.PHONG_DAO_TAO, password=MAT_KHAU
        )
        nganh, _ = NganhHoc.objects.get_or_create(ma_nganh='CNTT', defaults={'ten_nganh': 'Công nghệ thông tin'})
        self.tk_cv = TaiKhoan.objects.create_user(ten_dang_nhap='cv_cu', vai_tro=VaiTro.CO_VAN_HOC_TAP, password=MAT_KHAU)
        co_van = CoVanHocTap.objects.create(
            ma_co_van='C/V]]]', ho_ten='Cố vấn cũ', email='cvcu@test.local', so_dien_thoai='0913000001', ma_tai_khoan=self.tk_cv,
        )
        lop = LopHanhChinh.objects.create(ma_lop='LOPCU', ten_lop='Lớp cũ', ma_nganh=nganh, ma_co_van=co_van)
        self.tk_sv = TaiKhoan.objects.create_user(ten_dang_nhap='sv_xx', vai_tro=VaiTro.SINH_VIEN, password=MAT_KHAU)
        SinhVien.objects.create(
            ma_sinh_vien='///', ho_ten='Sinh viên mã lạ', ngay_sinh='2004-01-01', email='svcu@test.local',
            so_dien_thoai='0913000002', ma_lop=lop, ma_tai_khoan=self.tk_sv,
        )
        tk_gv = TaiKhoan.objects.create_user(ten_dang_nhap='gv_xx', vai_tro=VaiTro.GIANG_VIEN, password=MAT_KHAU)
        GiangVien.objects.create(
            ma_giang_vien='G/V', ho_ten='Giảng viên mã lạ', email='gvcu@test.local', so_dien_thoai='0913000003', ma_tai_khoan=tk_gv,
        )

    def test_trang_danh_sach_cua_phong_dao_tao_khong_sap(self):
        self.client.force_login(self.phong_dao_tao)
        for ten_url in ('people:sinh_vien_danh_sach', 'people:giang_vien_danh_sach', 'people:co_van_danh_sach'):
            response = self.client.get(reverse(ten_url))
            self.assertEqual(response.status_code, 200, ten_url)
            self.assertContains(response, 'Mã không hợp lệ')

    def test_trang_sinh_vien_phu_trach_va_trang_chu_co_van_khong_sap(self):
        self.client.force_login(self.tk_cv)
        response = self.client.get(reverse('enrollment:sinh_vien_phu_trach'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Sinh viên mã lạ')
        self.assertContains(response, 'Mã không hợp lệ')
        self.assertEqual(self.client.get(reverse('accounts:trang_chu')).status_code, 200)

    def test_sinh_vien_ma_la_vao_bang_diem_cua_toi_khong_sap(self):
        self.client.force_login(self.tk_sv)
        response = self.client.get(reverse('grades:bang_diem_cua_toi'), follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Mã sinh viên của bạn chứa ký tự không hợp lệ')

    def test_lenh_kiem_tra_ma_ho_so_bao_ma_cu_khong_hop_le(self):
        from io import StringIO

        from django.core.management import call_command
        from django.core.management.base import CommandError
        dau_ra = StringIO()
        with self.assertRaisesMessage(CommandError, 'Có 3 mã hồ sơ không hợp lệ'):
            call_command('kiem_tra_ma_ho_so', stdout=dau_ra)
        for doan in ("Sinh viên: mã '///'", "Giảng viên: mã 'G/V'", "Cố vấn học tập: mã 'C/V]]]'"):
            self.assertIn(doan, dau_ra.getvalue())


@tag('don_vi')
class LenhKiemTraMaHoSoSachTest(TestCase):
    def test_khong_co_ma_khong_hop_le(self):
        from io import StringIO

        from django.core.management import call_command
        tk = TaiKhoan.objects.create_user(ten_dang_nhap='gv_ok', vai_tro=VaiTro.GIANG_VIEN, password=MAT_KHAU)
        GiangVien.objects.create(ma_giang_vien='GV01', ho_ten='Hợp lệ', email='ok@test.local', so_dien_thoai='0915000001', ma_tai_khoan=tk)
        dau_ra = StringIO()
        call_command('kiem_tra_ma_ho_so', stdout=dau_ra)
        self.assertIn('Mọi mã hồ sơ đều hợp lệ', dau_ra.getvalue())


@tag('tich_hop')
class NhapExcelMaSinhVienKhongHopLeTest(TestCase):
    def setUp(self):
        self.phong_dao_tao = TaiKhoan.objects.create_user(
            ten_dang_nhap='pdt_excel_ma', vai_tro=VaiTro.PHONG_DAO_TAO, password=MAT_KHAU
        )
        self.lop = _tao_lop_hanh_chinh('CNTTEXMA')
        self.client.force_login(self.phong_dao_tao)

    def test_bao_loi_theo_so_dong_dong_hop_le_van_nhap(self):
        file_excel = _tao_file_excel([
            ['///', 'A', '2004-05-01', 'ma1@test.local', '0914000001', self.lop.pk, ''],
            ['SV 02', 'B', '2004-05-01', 'ma2@test.local', '0914000002', self.lop.pk, ''],
            ['SVOK03', 'C', '2004-05-01', 'ma3@test.local', '0914000003', self.lop.pk, ''],
            [2023004.0, 'D', '2004-05-01', 'ma4@test.local', '0914000004', self.lop.pk, ''],
        ])
        response = self.client.post(reverse('people:sinh_vien_nhap_excel'), {'file_excel': file_excel})
        ket_qua = response.context['ket_qua']
        self.assertEqual([(l['dong'], 'không hợp lệ' in l['loi']) for l in ket_qua['loi']], [(2, True), (3, True)])
        self.assertIn('Mã sinh viên "///"', ket_qua['loi'][0]['loi'])
        self.assertEqual(ket_qua['so_thanh_cong'], 2)
        self.assertTrue(SinhVien.objects.filter(pk='SVOK03').exists())
        self.assertTrue(SinhVien.objects.filter(pk='2023004').exists())
        self.assertFalse(SinhVien.objects.filter(pk='///').exists())


@tag('don_vi')
class LienHeKhongTrungServiceTest(TestCase):
    def setUp(self):
        self.lop = _tao_lop_hanh_chinh('LOPLH')  # cố vấn CV_LOPLH: cv@test.local, 0900000000
        tk = TaiKhoan.objects.create_user(ten_dang_nhap='gv_lh', vai_tro=VaiTro.GIANG_VIEN, password=MAT_KHAU)
        self.gv = GiangVien.objects.create(
            ma_giang_vien='GVLH', ho_ten='Giảng viên LH', email='GiangVien@Test.Local', so_dien_thoai='0912 345 678', ma_tai_khoan=tk,
        )

    def test_email_khong_phan_biet_hoa_thuong_sdt_so_sanh_sau_chuan_hoa(self):
        from .lien_he import chuan_hoa_so_dien_thoai, ho_so_trung_email, ho_so_trung_so_dien_thoai
        self.assertEqual(chuan_hoa_so_dien_thoai(' 0912.345-678 '), '0912345678')
        self.assertEqual(ho_so_trung_email('giangvien@test.local'), 'giảng viên GVLH (Giảng viên LH)')
        self.assertIn('cố vấn học tập CV_LOPLH', ho_so_trung_email('CV@TEST.LOCAL'))
        self.assertEqual(ho_so_trung_so_dien_thoai('0912345678'), 'giảng viên GVLH (Giảng viên LH)')
        self.assertIsNone(ho_so_trung_email('moi@test.local'))
        # Không tính chính hồ sơ đang sửa.
        self.assertIsNone(ho_so_trung_email('giangvien@test.local', bo_qua=self.gv))


@tag('tich_hop')
class LienHeKhongTrungFormTest(TestCase):
    def setUp(self):
        self.phong_dao_tao = TaiKhoan.objects.create_user(ten_dang_nhap='pdt_lh', vai_tro=VaiTro.PHONG_DAO_TAO, password=MAT_KHAU)
        self.lop = _tao_lop_hanh_chinh('LOPLHF')
        self.tk_gv = TaiKhoan.objects.create_user(ten_dang_nhap='gv_lhf', vai_tro=VaiTro.GIANG_VIEN, password=MAT_KHAU)
        self.gv = GiangVien.objects.create(
            ma_giang_vien='GVLHF', ho_ten='Giảng viên F', email='GiangVien@Test.Local', so_dien_thoai='0912 345 678', ma_tai_khoan=self.tk_gv,
        )
        self.client.force_login(self.phong_dao_tao)

    def _tao_sinh_vien(self, email, sdt, ma='SVLH1'):
        tk = TaiKhoan.objects.create_user(ten_dang_nhap=f'tk_{ma}', vai_tro=VaiTro.SINH_VIEN, password=MAT_KHAU)
        return self.client.post(reverse('people:sinh_vien_tao'), {
            'ma_sinh_vien': ma, 'ho_ten': 'SV', 'ngay_sinh': '2004-01-01', 'email': email,
            'so_dien_thoai': sdt, 'ma_lop': self.lop.pk, 'ma_tai_khoan': tk.pk,
        })

    def test_tao_sinh_vien_trung_email_giang_vien_khac_hoa_thuong_bi_bao_loi(self):
        response = self._tao_sinh_vien('giangvien@test.local', '0916000001')
        self.assertContains(response, 'đã được dùng cho giảng viên GVLHF')
        self.assertFalse(SinhVien.objects.filter(pk='SVLH1').exists())

    def test_tao_sinh_vien_trung_so_dien_thoai_khac_cach_viet_bi_bao_loi(self):
        response = self._tao_sinh_vien('svmoi@test.local', '0912345678')
        self.assertContains(response, 'Số điện thoại &quot;0912345678&quot; đã được dùng cho giảng viên GVLHF')
        self.assertFalse(SinhVien.objects.filter(pk='SVLH1').exists())

    def test_tao_giang_vien_trung_email_co_van_bi_bao_loi(self):
        tk = TaiKhoan.objects.create_user(ten_dang_nhap='gv_moi_lh', vai_tro=VaiTro.GIANG_VIEN, password=MAT_KHAU)
        response = self.client.post(reverse('people:giang_vien_tao'), {
            'ma_giang_vien': 'GVMOI', 'ho_ten': 'GV mới', 'email': 'CV@test.local', 'so_dien_thoai': '0916000002', 'ma_tai_khoan': tk.pk,
        })
        self.assertContains(response, 'đã được dùng cho cố vấn học tập')
        self.assertFalse(GiangVien.objects.filter(pk='GVMOI').exists())

    def test_sua_ho_so_giu_nguyen_email_sdt_cua_chinh_minh_van_luu_duoc(self):
        response = self.client.post(reverse('people:giang_vien_sua', args=['GVLHF']), {
            'ma_giang_vien': 'GVLHF', 'ho_ten': 'Giảng viên F đã sửa', 'email': 'giangvien@test.local',
            'so_dien_thoai': '0912345678', 'ma_tai_khoan': self.tk_gv.pk,
        })
        self.assertRedirects(response, reverse('people:giang_vien_danh_sach'))
        self.gv.refresh_from_db()
        self.assertEqual(self.gv.ho_ten, 'Giảng viên F đã sửa')

    def test_du_lieu_cu_dang_trung_khong_chan_sua_truong_khac(self):
        tk = TaiKhoan.objects.create_user(ten_dang_nhap='gv_cu_trung', vai_tro=VaiTro.GIANG_VIEN, password=MAT_KHAU)
        GiangVien.objects.create(ma_giang_vien='GVTRUNG', ho_ten='Trùng cũ', email='giangvien@test.local',
                                 so_dien_thoai='0912345678', ma_tai_khoan=tk)
        response = self.client.post(reverse('people:giang_vien_sua', args=['GVTRUNG']), {
            'ma_giang_vien': 'GVTRUNG', 'ho_ten': 'Đổi tên', 'email': 'giangvien@test.local',
            'so_dien_thoai': '0912345678', 'ma_tai_khoan': tk.pk,
        })
        self.assertRedirects(response, reverse('people:giang_vien_danh_sach'))

    def test_trang_thong_tin_ca_nhan_bao_trung_khong_lo_ho_so_khac(self):
        tk_sv = TaiKhoan.objects.create_user(ten_dang_nhap='sv_ttcn', vai_tro=VaiTro.SINH_VIEN, password=MAT_KHAU)
        sv = SinhVien.objects.create(ma_sinh_vien='SVTTCN', ho_ten='SV', ngay_sinh='2004-01-01', email='sv@test.local',
                                     so_dien_thoai='0916000003', ma_lop=self.lop, ma_tai_khoan=tk_sv)
        self.client.force_login(tk_sv)
        url = reverse('accounts:thong_tin_ca_nhan')
        response = self.client.post(url, {'email': 'GIANGVIEN@test.local', 'so_dien_thoai': '0912345678'})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Email này đã được dùng cho một hồ sơ khác')
        self.assertContains(response, 'Số điện thoại này đã được dùng cho một hồ sơ khác')
        self.assertNotContains(response, 'GVLHF')
        self.assertNotContains(response, 'Giảng viên F')
        sv.refresh_from_db()
        self.assertEqual(sv.email, 'sv@test.local')
        # Giữ nguyên email/sđt của chính mình thì lưu được.
        self.assertRedirects(self.client.post(url, {'email': 'SV@test.local', 'so_dien_thoai': '0916000003'}), url)


@tag('tich_hop')
class NhapExcelLienHeKhongTrungTest(TestCase):
    def setUp(self):
        self.client.force_login(TaiKhoan.objects.create_user(ten_dang_nhap='pdt_lhx', vai_tro=VaiTro.PHONG_DAO_TAO, password=MAT_KHAU))
        self.lop = _tao_lop_hanh_chinh('LOPLHX')  # cố vấn: cv@test.local, 0900000000

    def test_bao_loi_trung_du_lieu_co_san_va_trung_trong_tep(self):
        file_excel = _tao_file_excel([
            ['SVX01', 'A', '2004-05-01', 'CV@Test.Local', '0917000001', self.lop.pk, ''],
            ['SVX02', 'B', '2004-05-01', 'b@test.local', '0917000002', self.lop.pk, ''],
            ['SVX03', 'C', '2004-05-01', 'B@TEST.LOCAL', '0917 000 002', self.lop.pk, ''],
            ['SVX04', 'D', '2004-05-01', 'd@test.local', '0900000000', self.lop.pk, ''],
        ])
        response = self.client.post(reverse('people:sinh_vien_nhap_excel'), {'file_excel': file_excel})
        ket_qua = response.context['ket_qua']
        loi = {l['dong']: l['loi'] for l in ket_qua['loi']}
        self.assertEqual(sorted(loi), [2, 4, 5])
        self.assertIn('đã được dùng cho cố vấn học tập', loi[2])
        self.assertIn('trùng với dòng 3 trong tệp', loi[4])
        self.assertIn('Số điện thoại "0917000002" trùng với dòng 3', loi[4])  # đã chuẩn hóa từ "0917 000 002"
        self.assertIn('Số điện thoại "0900000000" đã được dùng cho', loi[5])
        self.assertEqual(ket_qua['so_thanh_cong'], 1)
        self.assertEqual(list(SinhVien.objects.values_list('pk', flat=True)), ['SVX02'])


def kiem_tra_dinh_dang_excel(test, noi_dung_tep):
    """Mọi sheet của file xuất: tiêu đề in đậm, cố định dòng 1, cột giãn theo nội dung dài nhất (tối đa 60)."""
    from io import BytesIO as _BytesIO

    from openpyxl import load_workbook as _load_workbook
    from openpyxl.utils import get_column_letter as _cot
    workbook = _load_workbook(_BytesIO(noi_dung_tep))
    for sheet in workbook.worksheets:
        test.assertEqual(sheet.freeze_panes, 'A2', sheet.title)
        test.assertTrue(all(cell.font.bold for cell in sheet[1] if cell.value is not None), sheet.title)
        for chi_so, cot in enumerate(sheet.iter_cols(values_only=True), start=1):
            dai_nhat = max(len(str(v)) for v in cot if v is not None)
            test.assertEqual(sheet.column_dimensions[_cot(chi_so)].width, min(max(dai_nhat + 2, 8), 60), (sheet.title, chi_so))
    return workbook


@tag('tich_hop')
class DinhDangFileMauNhapSinhVienTest(TestCase):
    def test_file_mau_gian_cot_theo_tieu_de_dai_in_dam_co_dinh(self):
        from openpyxl.utils import get_column_letter

        from .services import TIEU_DE_COT_MAU_SINH_VIEN
        self.client.force_login(TaiKhoan.objects.create_user(ten_dang_nhap='pdt_mau', vai_tro=VaiTro.PHONG_DAO_TAO, password=MAT_KHAU))
        response = self.client.get(reverse('people:sinh_vien_mau_excel'))
        self.assertEqual(response.status_code, 200)
        sheet = kiem_tra_dinh_dang_excel(self, response.content).active
        for chi_so, tieu_de in enumerate(TIEU_DE_COT_MAU_SINH_VIEN, start=1):
            self.assertEqual(sheet.column_dimensions[get_column_letter(chi_so)].width, min(max(len(tieu_de) + 2, 8), 60))


@tag('tich_hop')
class NhapExcelTenDangNhapKhongHopLeTest(TestCase):
    def setUp(self):
        self.client.force_login(TaiKhoan.objects.create_user(ten_dang_nhap='pdt_tdnx', vai_tro=VaiTro.PHONG_DAO_TAO, password=MAT_KHAU))
        self.lop = _tao_lop_hanh_chinh('LOPTDN')

    def test_cot_ten_dang_nhap_sai_quy_tac_bao_loi_theo_dong(self):
        file_excel = _tao_file_excel([
            ['SVTDN1', 'A', '2004-05-01', 'tdn1@test.local', '0918000001', self.lop.pk, '....'],
            ['SVTDN2', 'B', '2004-05-01', 'tdn2@test.local', '0918000002', self.lop.pk, 'sv_tdn2'],
            ['SVTDN3', 'C', '2004-05-01', 'tdn3@test.local', '0918000003', self.lop.pk, ''],
        ])
        ket_qua = self.client.post(reverse('people:sinh_vien_nhap_excel'), {'file_excel': file_excel}).context['ket_qua']
        self.assertEqual([l['dong'] for l in ket_qua['loi']], [2])
        self.assertIn('Tên đăng nhập "...." không hợp lệ', ket_qua['loi'][0]['loi'])
        self.assertEqual(sorted(t['ten_dang_nhap'] for t in ket_qua['tao_moi']), ['sv_tdn2', 'svtdn3'])


@tag('tich_hop')
class DinhDangEmailSoDienThoaiHoSoTest(TestCase):
    def setUp(self):
        self.client.force_login(TaiKhoan.objects.create_user(ten_dang_nhap='pdt_dd', vai_tro=VaiTro.PHONG_DAO_TAO, password=MAT_KHAU))
        self.lop = _tao_lop_hanh_chinh('LOPDD')

    def _tao_sinh_vien(self, ma='SVDD1', email='dd1@gmail.com', sdt='0919000001'):
        tk = TaiKhoan.objects.create_user(ten_dang_nhap=f'tk_{ma.lower()}_{TaiKhoan.objects.count()}', vai_tro=VaiTro.SINH_VIEN, password=MAT_KHAU)
        return self.client.post(reverse('people:sinh_vien_tao'), {
            'ma_sinh_vien': ma, 'ho_ten': 'SV', 'ngay_sinh': '2004-01-01', 'email': email,
            'so_dien_thoai': sdt, 'ma_lop': self.lop.pk, 'ma_tai_khoan': tk.pk,
        })

    def test_email_sai_dinh_dang_bi_chan(self):
        for email in ('abc', 'gmail', 'abc@', '@gmail.com', 'a b@gmail.com'):
            response = self._tao_sinh_vien(email=email)
            self.assertContains(response, 'Email không đúng định dạng', msg_prefix=email)
        self.assertFalse(SinhVien.objects.exists())

    def test_so_dien_thoai_sai_dinh_dang_bi_chan_dung_thi_tu_chuan_hoa(self):
        for sdt in ('12ab', '091234567', '1912345678', '091234567890'):
            response = self._tao_sinh_vien(sdt=sdt)
            self.assertContains(response, 'Số điện thoại phải gồm 10 chữ số', msg_prefix=sdt)
        self.assertRedirects(self._tao_sinh_vien(sdt='0919 000.001'), reverse('people:sinh_vien_danh_sach'))
        self.assertEqual(SinhVien.objects.get(pk='SVDD1').so_dien_thoai, '0919000001')

    def test_trung_ma_khac_hoa_thuong_bao_loi_tieng_viet(self):
        self._tao_sinh_vien(ma='SVDD1')
        response = self._tao_sinh_vien(ma='svdd1', email='dd2@gmail.com', sdt='0919000002')
        self.assertContains(response, 'Mã sinh viên &quot;svdd1&quot; đã tồn tại.')
        self.assertNotContains(response, 'Ma sinh vien')
        self.assertEqual(SinhVien.objects.count(), 1)

    def test_du_lieu_cu_sai_dinh_dang_van_sua_duoc_truong_khac(self):
        tk = TaiKhoan.objects.create_user(ten_dang_nhap='gv_cu_dd', vai_tro=VaiTro.GIANG_VIEN, password=MAT_KHAU)
        GiangVien.objects.create(ma_giang_vien='GVDD', ho_ten='Cũ', email='khong-phai-email', so_dien_thoai='12ab', ma_tai_khoan=tk)
        response = self.client.post(reverse('people:giang_vien_sua', args=['GVDD']), {
            'ma_giang_vien': 'GVDD', 'ho_ten': 'Đã sửa tên', 'email': 'khong-phai-email', 'so_dien_thoai': '12ab', 'ma_tai_khoan': tk.pk,
        })
        self.assertRedirects(response, reverse('people:giang_vien_danh_sach'))
        # Đổi email sang giá trị sai định dạng khác thì bị chặn.
        response = self.client.post(reverse('people:giang_vien_sua', args=['GVDD']), {
            'ma_giang_vien': 'GVDD', 'ho_ten': 'Đã sửa tên', 'email': 'van-sai', 'so_dien_thoai': '12ab', 'ma_tai_khoan': tk.pk,
        })
        self.assertContains(response, 'Email không đúng định dạng')


@tag('tich_hop')
class NhapExcelDinhDangLienHeTest(TestCase):
    def setUp(self):
        self.client.force_login(TaiKhoan.objects.create_user(ten_dang_nhap='pdt_ddx', vai_tro=VaiTro.PHONG_DAO_TAO, password=MAT_KHAU))
        self.lop = _tao_lop_hanh_chinh('LOPDDX')

    def test_email_so_dien_thoai_sai_bao_loi_o_so_bi_mat_so_0_tu_them_lai(self):
        file_excel = _tao_file_excel([
            ['SVDX1', 'A', '2004-05-01', 'gmail', '0919100001', self.lop.pk, ''],
            ['SVDX2', 'B', '2004-05-01', 'b@gmail.com', '12ab', self.lop.pk, ''],
            ['SVDX3', 'C', '2004-05-01', 'c@gmail.com', 919100003, self.lop.pk, ''],
            ['SVDX4', 'D', '2004-05-01', 'd@gmail.com', '0919 100 004', self.lop.pk, ''],
        ])
        ket_qua = self.client.post(reverse('people:sinh_vien_nhap_excel'), {'file_excel': file_excel}).context['ket_qua']
        loi = {l['dong']: l['loi'] for l in ket_qua['loi']}
        self.assertEqual(sorted(loi), [2, 3])
        self.assertIn('Email "gmail" không đúng định dạng', loi[2])
        self.assertIn('Số điện thoại "12ab" không hợp lệ', loi[3])
        self.assertEqual(SinhVien.objects.get(pk='SVDX3').so_dien_thoai, '0919100003')
        self.assertEqual(SinhVien.objects.get(pk='SVDX4').so_dien_thoai, '0919100004')
