from datetime import date, datetime
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# Độ rộng cột (đơn vị ký tự của Excel): theo nội dung dài nhất trong cột, cộng lề, trong khoảng [tối thiểu, tối đa].
DO_RONG_COT_TOI_THIEU = 8
DO_RONG_COT_TOI_DA = 60
LE_COT = 2


def _an_toan(gia_tri):
    if isinstance(gia_tri, str):
        return ILLEGAL_CHARACTERS_RE.sub('', gia_tri)
    return gia_tri


def _do_dai_hien_thi(gia_tri):
    if gia_tri is None:
        return 0
    if isinstance(gia_tri, datetime):
        gia_tri = gia_tri.strftime('%Y-%m-%d %H:%M')
    elif isinstance(gia_tri, date):
        gia_tri = gia_tri.isoformat()
    return max((len(dong) for dong in str(gia_tri).splitlines()), default=0)


def dinh_dang_sheet(sheet):
    """Dòng 1 là tiêu đề: in đậm, cố định (freeze panes) khi cuộn; mỗi cột giãn theo nội dung dài nhất
    (tối đa DO_RONG_COT_TOI_DA ký tự). Gọi lại nhiều lần vẫn cho cùng kết quả."""
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = 'A2'
    for chi_so, cot in enumerate(sheet.iter_cols(values_only=True), start=1):
        dai_nhat = max((_do_dai_hien_thi(gia_tri) for gia_tri in cot), default=0)
        sheet.column_dimensions[get_column_letter(chi_so)].width = min(
            max(dai_nhat + LE_COT, DO_RONG_COT_TOI_THIEU), DO_RONG_COT_TOI_DA,
        )


def tao_workbook_excel(ten_sheet, tieu_de_cot, cac_dong):
    """Tạo một workbook một sheet, dòng đầu là tiêu đề cột in đậm — dùng chung cho mọi
    chức năng xuất Excel trong dự án (không phụ thuộc app nghiệp vụ nào)."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = ten_sheet
    sheet.append(list(tieu_de_cot))
    for dong in cac_dong:
        sheet.append([_an_toan(gia_tri) for gia_tri in dong])
    dinh_dang_sheet(sheet)
    return workbook


def phan_hoi_excel(workbook, ten_file):
    """Đưa workbook vào HttpResponse tải về dạng .xlsx. Định dạng lại mọi sheet (kể cả sheet view tự thêm sau
    tao_workbook_excel) để file xuất nào cũng có tiêu đề in đậm, cố định và cột vừa nội dung."""
    for sheet in workbook.worksheets:
        dinh_dang_sheet(sheet)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{ten_file}"'
    return response
